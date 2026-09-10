# -*- coding: utf-8 -*-
"""DOCX report assembly helpers (python-docx) for the report MCP."""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt
from docx.oxml.ns import qn  # for CJK font

# Reuse the chart helper's path sandbox for consistency.
from .charts import _resolve_allowed_path


def _set_cjk_font(run, font_name: str = "宋体") -> None:
    """Apply a CJK font name so Chinese text doesn't fall back to tofu."""
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        from docx.oxml import OxmlElement

        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)


def _apply_section_paragraph(
    doc: Document,
    text: str,
    style: str = "Normal",
    font_name: str = "宋体",
    font_size: float = 11.0,
) -> None:
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    _set_cjk_font(run, font_name)


def _add_heading(doc: Document, text: str, level: int) -> None:
    size_map = {1: 18.0, 2: 14.0, 3: 12.0}
    style_name = f"Heading {max(1, min(level, 9))}"
    p = doc.add_paragraph(style=style_name)
    run = p.add_run(text)
    run.font.size = Pt(size_map.get(level, 11.0))
    run.bold = True
    _set_cjk_font(run, "黑体")


def _add_table(doc: Document, table_payload: dict[str, Any]) -> None:
    """Insert a table from ``{headers: [...], rows: [[...], ...]}``."""
    headers = list(table_payload.get("headers") or [])
    rows = list(table_payload.get("rows") or [])
    if not headers and not rows:
        return
    if not headers and rows:
        headers = [f"col_{i}" for i in range(len(rows[0]))]
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = str(h)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            _set_cjk_font(run, "黑体")
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            if j >= len(headers):
                break
            cell = table.rows[i].cells[j]
            cell.text = "" if val is None else str(val)
            for run in cell.paragraphs[0].runs:
                _set_cjk_font(run, "宋体")


def _insert_chart(doc: Document, chart_path: str) -> bool:
    """Insert a PNG chart into the document. Returns True on success."""
    if not chart_path:
        return False
    p_path, err = _resolve_allowed_path(chart_path)
    if err is not None or p_path is None or not p_path.is_file():
        return False
    try:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(p_path), width=Inches(5.5))
        return True
    except Exception:  # noqa: BLE001
        return False


def _set_header_footer(doc: Document, header: str, footer: str) -> None:
    for section in doc.sections:
        if header:
            h = section.header.paragraphs[0]
            h.text = header
            for run in h.runs:
                run.font.size = Pt(9)
                _set_cjk_font(run, "宋体")
        if footer:
            f = section.footer.paragraphs[0]
            f.text = footer
            for run in f.runs:
                run.font.size = Pt(9)
                _set_cjk_font(run, "宋体")


def _set_page_margins(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)


def render_docx_report(
    title: str,
    sections: list[dict[str, Any]],
    output_path: str,
    header: str = "",
    footer: str = "",
    style: str = "default",
) -> dict:
    """Build a DOCX report and write it to ``output_path``.

    Args:
        title: report title (H1).
        sections: list of dicts with keys
            ``heading`` (str, optional), ``level`` (int 1-3, default 2),
            ``paragraphs`` (list[str]), ``table`` (dict with headers/rows),
            ``charts`` (list[str] of PNG paths).
        output_path: absolute path inside the workspace sandbox.
        header/footer: optional strings.
        style: reserved for future templates; ignored for now.

    Returns:
        dict with ``ok, output_path, sections_count, charts_inserted, message``.
    """
    if not output_path:
        return _err("output_path 不能为空 / output_path is required")
    if not str(output_path).lower().endswith(".docx"):
        return _err("output_path 必须以 .docx 结尾 / output_path must end with .docx")

    p_path, err = _resolve_allowed_path(output_path)
    if err is not None or p_path is None:
        return _err((err or {"message": "路径解析失败"})["message"])

    p_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _set_page_margins(doc)
    _set_header_footer(doc, header, footer)

    # Title
    _add_heading(doc, title or "未命名报告 / Untitled Report", level=1)

    charts_inserted = 0
    for sec in sections or []:
        level = int(sec.get("level") or 2)
        heading = sec.get("heading")
        if heading:
            _add_heading(doc, str(heading), level)
        for para in sec.get("paragraphs") or []:
            _apply_section_paragraph(doc, str(para), style="Normal")
        if sec.get("table"):
            _add_table(doc, sec["table"])
        for chart_path in sec.get("charts") or []:
            if _insert_chart(doc, str(chart_path)):
                charts_inserted += 1

    try:
        doc.save(str(p_path))
    except Exception as exc:  # noqa: BLE001
        return _err(f"保存 docx 失败 / Failed to save docx: {exc}")

    return {
        "ok": True,
        "output_path": str(p_path),
        "sections_count": len(sections or []),
        "charts_inserted": charts_inserted,
        "message": f"已生成报告 / Report generated: {p_path.name} "
        f"(sections={len(sections or [])}, charts_inserted={charts_inserted})",
    }


def _err(message: str) -> dict:
    return {"ok": False, "output_path": "", "sections_count": 0, "charts_inserted": 0, "message": message}
