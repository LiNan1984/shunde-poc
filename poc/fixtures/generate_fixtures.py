# -*- coding: utf-8 -*-
"""Generate demo fixture files under poc/fixtures/ for Excel guard POC."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

FIXTURES_DIR = Path(__file__).resolve().parent
KB_FIXTURES_DIR = FIXTURES_DIR / "kb"

# doc_id -> (filename, builder name in poc.kb_mcp.fixtures / this module)
KB_PDF_BUILDERS: dict[str, tuple[str, str]] = {
    "credit-policy": ("信贷政策.pdf", "write_native_text_pdf"),
    "pricing": ("产品定价.pdf", "write_table_image_pdf"),
    "branch-balance": ("网点余额图.pdf", "write_branch_balance_chart_pdf"),
}

# Exact figures for the chart PDF. They exist ONLY as bar heights: the body
# text must never contain them, so image-modality golden QA items can only be
# answered by a vision model (analyze_page with --live).
BRANCH_BALANCE_VALUES: dict[str, float] = {
    "BranchA": 1.23,
    "BranchB": 2.46,
    "BranchC": 3.69,
}


def _ensure_repo_on_path() -> None:
    repo_root = str(Path(__file__).resolve().parents[2])
    if repo_root not in sys.path:
        sys.path.insert(0, repo_root)


def write_corrupt_xlsx(path: Path) -> None:
    path.write_bytes(b"NOT_A_VALID_XLSX_OR_ZIP_PACKAGE_FOR_POC_DEMO")


def write_encoding_latin1_csv(path: Path) -> None:
    # Latin-1 content with non-UTF-8 bytes (é, ü)
    text = "id,name,city\n1,Jos\xe9,M\xfcnchen\n2,Fran\xe7ois,Caf\xe9\n"
    path.write_bytes(text.encode("latin-1"))


def write_shunde_income_xlsx(
    dest: Path,
    source: Path | None = None,
) -> None:
    """Copy the 收支明细 workbook and replace 青岛 with 顺德 (keep as a demo file)."""
    from openpyxl import load_workbook

    repo_root = Path(__file__).resolve().parents[2]
    src = source or repo_root / "收支明细测试数据_v5_100覆盖(1).xlsx"
    wb = load_workbook(src)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "青岛" in cell.value:
                    cell.value = cell.value.replace("青岛", "顺德")
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


def write_large_chunk_demo_xlsx(path: Path, rows: int = 6000) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "DemoData"
    ws.append(["row_id", "amount", "note"])
    for i in range(1, rows + 1):
        ws.append([i, i * 10.5, f"row-{i}"])
    wb.save(path)


def write_branch_balance_chart_pdf(path: Path) -> Path:
    """Deterministic one-page chart PDF (网点余额图): answers live in bar heights.

    matplotlib draws the bars; the page keeps only a CJK title plus the image,
    so the exact figures never appear as text and offline retrieval cannot
    recall them (quantified by scripts/kb_eval.py, image group). English chart
    labels avoid depending on a CJK font for matplotlib; the PDF title uses
    reportlab's built-in STSong-Light (same as poc/kb_mcp/fixtures.py).
    """
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Image as RLImage
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    _ensure_repo_on_path()
    from poc.kb_mcp.fixtures import _cjk_style

    path.parent.mkdir(parents=True, exist_ok=True)
    png = path.with_suffix(".png")
    names = list(BRANCH_BALANCE_VALUES)
    values = [BRANCH_BALANCE_VALUES[n] for n in names]
    fig, ax = plt.subplots(figsize=(7.0, 4.0), dpi=150)
    ax.bar(names, values, color="#4472c4")
    ax.set_title("Branch Loan Balance (Q4)")
    ax.set_xlabel("Branch")
    ax.set_ylabel("Balance (100M CNY)")
    ax.set_ylim(0, 4.0)
    # Round ticks only: a tick label must never equal one of the exact values.
    ax.set_yticks([0, 0.5, 1.0, 1.5, 2.0, 2.5, 3.0, 3.5, 4.0])
    fig.tight_layout()
    fig.savefig(str(png))
    plt.close(fig)

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    story = [
        Paragraph("各网点贷款余额柱状图（单位：亿元）", _cjk_style(14)),
        Spacer(1, 10),
        RLImage(str(png), width=150 * mm, height=85 * mm),
    ]
    SimpleDocTemplate(str(path), pagesize=A4).build(story)
    return path


def write_kb_pdf(doc_id: str, dest_dir: Path | None = None) -> Path:
    """Build the deterministic KB fixture PDF for ``doc_id`` into ``dest_dir``."""
    _ensure_repo_on_path()
    from poc.kb_mcp.fixtures import write_native_text_pdf, write_table_image_pdf

    try:
        filename, builder_name = KB_PDF_BUILDERS[doc_id]
    except KeyError as exc:
        raise KeyError(f"unknown KB fixture doc_id: {doc_id}") from exc
    dest = (dest_dir or KB_FIXTURES_DIR) / filename
    builder = {
        "write_native_text_pdf": write_native_text_pdf,
        "write_table_image_pdf": write_table_image_pdf,
        "write_branch_balance_chart_pdf": write_branch_balance_chart_pdf,
    }[builder_name]
    return builder(dest)


def ensure_kb_fixtures(dest_dir: Path | None = None) -> dict[str, Path]:
    """Regenerate every KB fixture PDF deterministically; return doc_id -> path."""
    return {doc_id: write_kb_pdf(doc_id, dest_dir) for doc_id in KB_PDF_BUILDERS}


def main(kb_only: bool = False) -> None:
    if not kb_only:
        FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
        corrupt = FIXTURES_DIR / "corrupt.xlsx"
        encoding_csv = FIXTURES_DIR / "encoding_latin1.csv"
        large = FIXTURES_DIR / "large_chunk_demo.xlsx"

        write_corrupt_xlsx(corrupt)
        write_encoding_latin1_csv(encoding_csv)
        write_large_chunk_demo_xlsx(large, rows=6000)
        print(f"Wrote {corrupt}")
        print(f"Wrote {encoding_csv}")
        print(f"Wrote {large}")

        shunde_income = FIXTURES_DIR / "收支明细_顺德.xlsx"
        src = Path(__file__).resolve().parents[2] / "收支明细测试数据_v5_100覆盖(1).xlsx"
        if src.is_file():
            write_shunde_income_xlsx(shunde_income, source=src)
            print(f"Wrote {shunde_income}")
        elif shunde_income.is_file():
            print(f"keep existing {shunde_income}")
        else:
            print(f"skip {shunde_income.name}: source missing")

    for doc_id, path in ensure_kb_fixtures().items():
        print(f"Wrote {path} (doc_id={doc_id})")


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--kb-only",
        action="store_true",
        help="only (re)generate the KB fixture PDFs under poc/fixtures/kb/",
    )
    main(kb_only=parser.parse_args().kb_only)
