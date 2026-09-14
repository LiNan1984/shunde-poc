# -*- coding: utf-8 -*-
"""L1 read/write surface for spreadsheet QA: compressed describe, markdown slices, audit.

Design borrowed from the Shortcut spreadsheet-agent writeup: reads are a
compression act (header voting, formula aliasing, merge context) and writes
get a classified self-check (formula errors surfaced, suspicious constants
flagged) instead of an unreviewable wall of diffs.
"""

from __future__ import annotations

import csv
import logging
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from .guards import (
    _TEXT_SUFFIXES,
    _XLSX_SUFFIXES,
    _check_file_size,
    _resolve_allowed_path,
    _safe_name,
)

logger = logging.getLogger("poc.excel_guard")

_MD_CELL_MAX_CHARS = 40
_MD_MAX_COLS = 256
_MD_MAX_SHEETS = 200
_HEADER_SCAN_ROWS = 5
_FORMULA_SCAN_ROWS = 200
_ALIAS_MIN_COUNT = 3
_SLICE_MAX_ROWS = 2000
_AUDIT_MAX_ROWS = 100_000
_AUDIT_MAX_ENTRIES = 50

_EXCEL_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!"}
_FORMULA_LITERAL_RE = re.compile(r"(?<![A-Z0-9$.])\d{3,}(?!\d)")
_UNCACHED_PREVIEW_NOTE = (
    "公式列未缓存计算值（openpyxl 生成的新文件），预览中以 =公式 形式显示。"
)

DEFAULT_PREVIEW_ROWS = 10
MAX_PREVIEW_ROWS = 100


# --- shared helpers ------------------------------------------------------------


def _md_cell(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("|", "\\|").replace("\r", " ").replace("\n", " ")
    if len(text) > _MD_CELL_MAX_CHARS:
        text = text[:_MD_CELL_MAX_CHARS] + "…"
    return text


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _trim_trailing_empty(row: list[Any]) -> list[Any]:
    values = list(row)
    while values and values[-1] in (None, ""):
        values.pop()
    return values


def _used_cols(ws: Any) -> int:
    n = int(getattr(ws, "max_column", 0) or 0)
    if n <= 0:
        return 0
    return min(n, _MD_MAX_COLS)


def _iter_kwargs(ws: Any, **extra: Any) -> dict[str, Any]:
    kwargs = dict(extra)
    cols = _used_cols(ws)
    if cols:
        kwargs["max_col"] = cols
    return kwargs


def _denied(err: dict | None, **extra: Any) -> dict[str, Any]:
    """Sandbox denials must not echo the outside path/filename."""
    err = err or {"ok": False, "issue": "path_denied", "message": "路径无效 / Invalid path"}
    issue = err.get("issue") or "path_denied"
    if issue == "path_denied":
        message = "路径越界 / Path outside workspace"
    else:
        message = err.get("message") or "请求失败 / request failed"
    out: dict[str, Any] = {
        "ok": False,
        "issue": issue,
        "kind": "",
        "sheets": [],
        "markdown": "",
        "message": message,
    }
    out.update(extra)
    return out


def _rows_to_markdown(rows: list[list[Any]]) -> str:
    """Render rows as a markdown table; the first physical row is the header."""
    cleaned = [[_md_cell(c) for c in _trim_trailing_empty(list(row))] for row in rows]
    width = max((len(r) for r in cleaned), default=0)
    if width == 0:
        return ""
    for row in cleaned:
        row.extend([""] * (width - len(row)))
    parts = [
        "| " + " | ".join(cleaned[0]) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    parts.extend("| " + " | ".join(row) + " |" for row in cleaned[1:])
    return "\n".join(parts)


def _pick_header_row(rows: list[list[Any]]) -> int:
    """Vote the most table-like row among the first few as the header row.

    Heuristic: prefer rows whose non-empty cells are mostly text, then rows
    that fill the most columns. Returns a 0-based index into ``rows``.
    """
    best_idx, best_score = 0, -1.0
    for i, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        cells = [c for c in row if c not in (None, "")]
        if not cells:
            continue
        str_ratio = sum(
            1 for c in cells if isinstance(c, str) and not c.startswith("=")
        ) / len(cells)
        fill = len(cells) / max(1, len(row))
        score = str_ratio * 2 + fill
        if score > best_score:
            best_idx, best_score = i, score
    return best_idx


def _resolve_spreadsheet(path: str) -> tuple[Path | None, str, dict | None]:
    """Sandbox-check a path and classify it as xlsx / csv / unsupported.

    Denial messages never echo the raw path (prompt-injection hygiene):
    the workspace-relative name is fine, the absolute path is not.
    """
    file_path, err = _resolve_allowed_path(path)
    if err is not None or file_path is None:
        err = err or {"ok": False, "issue": "path_denied", "message": "路径无效 / Invalid path"}
        if err.get("issue") == "path_denied":
            err = {
                "ok": False,
                "issue": "path_denied",
                "message": "路径越界 / Path outside workspace",
            }
        return None, "", err
    if not file_path.exists():
        return None, "", {
            "ok": False,
            "issue": "missing",
            "message": f"文件不存在 / File not found: {_safe_name(file_path)}",
        }
    size_err = _check_file_size(file_path)
    if size_err is not None:
        return None, "", size_err
    suffix = file_path.suffix.lower()
    if suffix in _XLSX_SUFFIXES:
        return file_path, "xlsx", None
    if suffix in _TEXT_SUFFIXES:
        return file_path, "csv", None
    return None, "", {
        "ok": False,
        "message": (
            f"不支持的格式 / Unsupported format: {_safe_name(file_path)}。"
            "请使用 .xlsx / .xlsm / .csv 等支持的格式。"
        ),
    }


def _clamp_preview(preview_rows: int) -> int:
    if preview_rows < 0:
        return 0
    return min(preview_rows, MAX_PREVIEW_ROWS)


def _count_text_lines(file_path: Path) -> int:
    """Stream-count lines in constant memory (last unterminated fragment counts)."""
    newline_count = 0
    last_byte = -1
    with open(file_path, "rb") as fh:
        while True:
            buf = fh.read(1 << 20)
            if not buf:
                break
            newline_count += buf.count(b"\n")
            last_byte = buf[-1]
    return newline_count + (1 if last_byte not in (-1, 0x0A) else 0)


# --- xlsx internals ------------------------------------------------------------


def _merged_ranges_by_sheet(xlsx_path: Path) -> dict[str, list[str]]:
    """Sheet name -> merged range strings, parsed straight from the xlsx zip.

    read_only worksheets do not expose merged_cells, and a full non-read-only
    load would defeat the streaming guards — so read workbook.xml + rels +
    each sheet's mergeCells elements directly. Best-effort: any parse failure
    returns {} and describe simply omits merge context.
    """
    try:
        with zipfile.ZipFile(xlsx_path) as zf:
            ns_main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            ns_rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
            ns_pkg_rel = "{http://schemas.openxmlformats.org/package/2006/relationships}"
            wb_xml = ElementTree.fromstring(zf.read("xl/workbook.xml"))
            rels_xml = ElementTree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
            targets = {
                rel.get("Id"): rel.get("Target") or ""
                for rel in rels_xml.findall(f"{ns_pkg_rel}Relationship")
            }
            merged: dict[str, list[str]] = {}
            sheets_el = wb_xml.find(f"{ns_main}sheets")
            for sheet in sheets_el if sheets_el is not None else []:
                name = sheet.get("name") or ""
                rid = sheet.get(f"{ns_rel}id") or ""
                target = targets.get(rid, "")
                if target.startswith("/"):
                    xml_path = target.lstrip("/")
                elif target.startswith("xl/"):
                    xml_path = target
                else:
                    xml_path = "xl/" + target
                try:
                    sheet_xml = ElementTree.fromstring(zf.read(xml_path))
                except KeyError:
                    continue
                ranges = [
                    block.get("ref")
                    for block in sheet_xml.iter(f"{ns_main}mergeCell")
                    if block.get("ref")
                ]
                if ranges:
                    merged[name] = ranges[:100]
            return merged
    except Exception as exc:  # noqa: BLE001 — best-effort context
        logger.warning("mergeCells scan failed: %s", exc)
        return {}


def _cell_ref(row_idx: int, col_idx: int) -> str:
    return f"{get_column_letter(col_idx)}{row_idx}"


def _alias_formulas_in_column(column_formulas: list[tuple[str, str]]) -> dict[str, int]:
    """Count formulas in one column normalized to row 1 of that column.

    =A2*B2 and =A3*B3 both become =A1*B1, so a column of identical formulas
    collapses to one key: {normalized: count}.
    """
    counts: dict[str, int] = {}
    for cell_ref, formula in column_formulas:
        try:
            normalized = Translator(formula, origin=cell_ref).translate_formula(
                re.sub(r"\d+", "1", cell_ref, count=1)
            )
        except Exception:  # noqa: BLE001 — cross-sheet / odd refs stay verbatim
            normalized = formula
        counts[normalized] = counts.get(normalized, 0) + 1
    return counts


def _scan_formula_patterns(ws: Any, max_row: int) -> list[dict[str, Any]]:
    """Scan the top rows in formula mode and return aliased repeated patterns."""
    by_col: dict[int, list[tuple[str, str]]] = {}
    for row_idx, row in enumerate(
        ws.iter_rows(
            max_row=min(_FORMULA_SCAN_ROWS, max_row or 1),
            **_iter_kwargs(ws, values_only=True),
        ),
        start=1,
    ):
        for col_idx, value in enumerate(row, start=1):
            if isinstance(value, str) and value.startswith("="):
                by_col.setdefault(col_idx, []).append(
                    (_cell_ref(row_idx, col_idx), value)
                )
    aliases: list[dict[str, Any]] = []
    for col_idx in sorted(by_col):
        counts = _alias_formulas_in_column(by_col[col_idx])
        rank = 0
        for normalized, count in sorted(
            counts.items(), key=lambda kv: (-kv[1], kv[0])
        ):
            if count < _ALIAS_MIN_COUNT:
                continue
            rank += 1
            aliases.append(
                {
                    "alias": f"{get_column_letter(col_idx)}-F{rank}",
                    "column": get_column_letter(col_idx),
                    "pattern": normalized,
                    "count": count,
                    "scanned_rows": min(_FORMULA_SCAN_ROWS, max_row),
                }
            )
            if rank >= 5:
                break
    return aliases


def _formula_map(ws: Any, max_row: int) -> dict[str, str]:
    """cell ref -> formula for the scanned region (used to fill blank previews)."""
    out: dict[str, str] = {}
    for row_idx, row in enumerate(
        ws.iter_rows(
            max_row=min(_FORMULA_SCAN_ROWS, max_row or 1),
            **_iter_kwargs(ws, values_only=True),
        ),
        start=1,
    ):
        for col_idx, value in enumerate(row, start=1):
            if isinstance(value, str) and value.startswith("="):
                out[_cell_ref(row_idx, col_idx)] = value
    return out


def _describe_xlsx(xlsx_path: Path, preview_rows: int) -> list[dict[str, Any]]:
    merged = _merged_ranges_by_sheet(xlsx_path)
    wb_v = load_workbook(xlsx_path, read_only=True, data_only=True)
    wb_f = load_workbook(xlsx_path, read_only=True, data_only=False)
    sheets: list[dict[str, Any]] = []
    try:
        names = wb_v.sheetnames[:_MD_MAX_SHEETS]
        for name in names:
            ws = wb_v[name]
            ws_f = wb_f[name]
            max_row = int(ws.max_row or 0)
            max_col = int(ws.max_column or 0)
            scan_kwargs = _iter_kwargs(ws, values_only=True)
            header_rows = [
                list(row)
                for row in ws.iter_rows(
                    max_row=min(_HEADER_SCAN_ROWS, max_row or 1), **scan_kwargs
                )
            ]
            header_idx = _pick_header_row(header_rows) if header_rows else 0
            header = (
                [_cell_str(c) for c in _trim_trailing_empty(header_rows[header_idx])]
                if header_rows
                else []
            )
            preview_limit = min(max_row, 1 + max(0, preview_rows)) if max_row else 0
            preview: list[list[Any]] = []
            if preview_limit:
                preview = [
                    list(row)
                    for row in ws.iter_rows(max_row=preview_limit, **scan_kwargs)
                ]
            if preview:
                formulas = _formula_map(ws_f, max_row)
                uncached = False
                for r_idx, row in enumerate(preview, start=1):
                    for c_idx in range(len(row)):
                        if row[c_idx] is None:
                            ref = _cell_ref(r_idx, c_idx + 1)
                            if ref in formulas:
                                row[c_idx] = formulas[ref]
                                uncached = True
                note = _UNCACHED_PREVIEW_NOTE if uncached else ""
            else:
                note = ""
            preview = [_trim_trailing_empty(row) for row in preview]
            preview_str = [[_cell_str(c) for c in row] for row in preview]
            sheets.append(
                {
                    "name": name,
                    "row_count": max_row,
                    "column_count": max_col,
                    "headers": header,
                    "preview": preview_str,
                    "rows": max_row,
                    "cols": max_col,
                    "header_row": header_idx + 1,
                    "header": header,
                    "merged_ranges": merged.get(name) or [],
                    "formula_aliases": _scan_formula_patterns(ws_f, max_row),
                    "markdown": _rows_to_markdown(preview) if preview else "",
                    "note": note,
                }
            )
    finally:
        wb_v.close()
        wb_f.close()
    return sheets


# --- describe_workbook ---------------------------------------------------------


def describe_workbook(path: str, preview_rows: int = DEFAULT_PREVIEW_ROWS) -> dict:
    """Compressed workbook structure: sheets, dims, voted header, merges,
    formula aliases, markdown preview. One call replaces blind pandas probing.
    """
    file_path, kind, err = _resolve_spreadsheet(path)
    if err is not None or file_path is None:
        return _denied(err)
    preview_rows = _clamp_preview(preview_rows)

    if kind == "csv":
        from .guards import detect_encoding

        enc = detect_encoding(str(file_path))
        encoding = enc.get("encoding") or "utf-8"
        if encoding in {"unknown", "ascii"}:
            encoding = "utf-8" if encoding == "ascii" else encoding
        total = _count_text_lines(file_path)
        rows: list[list[Any]] = []
        read_encoding = "utf-8" if encoding in {"unknown"} else encoding
        take = 1 + preview_rows
        try:
            with open(file_path, "r", encoding=read_encoding, errors="replace", newline="") as fh:
                for i, row in enumerate(csv.reader(fh)):
                    if i >= max(take, _HEADER_SCAN_ROWS):
                        break
                    rows.append(list(row))
        except OSError as exc:
            return {
                "ok": False,
                "issue": "unreadable",
                "kind": "csv",
                "sheets": [],
                "markdown": "",
                "message": f"无法读取文件 / Cannot read file: {exc}",
            }
        header_idx = _pick_header_row(rows) if rows else 0
        header = list(rows[header_idx]) if rows else []
        preview = rows[:take]
        sheet_rec = {
            "name": file_path.stem,
            "row_count": total,
            "column_count": len(header),
            "headers": header,
            "preview": preview,
            "rows": total,
            "cols": len(header),
            "header_row": header_idx + 1,
            "header": header,
            "merged_ranges": [],
            "formula_aliases": [],
            "markdown": _rows_to_markdown(preview) if preview else "",
            "note": f"encoding={read_encoding} (confidence={enc.get('confidence', 0):.2f})",
        }
        return {
            "ok": True,
            "kind": "csv",
            "path": str(file_path),
            "sheets": [sheet_rec],
            "message": (
                f"CSV {total} 行；表头在第 {header_idx + 1} 行，"
                f"编码 {read_encoding}。预览 {len(preview)} 行。"
            ),
        }

    try:
        sheets = _describe_xlsx(file_path, preview_rows)
    except MemoryError:
        return {
            "ok": False,
            "kind": "xlsx",
            "sheets": [],
            "message": "文件解析内存不足 / Out of memory while parsing workbook",
        }
    except Exception as exc:  # noqa: BLE001 — surface as corrupt-likely
        return {
            "ok": False,
            "kind": "xlsx",
            "sheets": [],
            "message": (
                f"无法解析工作簿 / Cannot parse workbook: {type(exc).__name__}: {exc}。"
                "请先调用 detect_corrupt_workbook。"
            ),
        }
    names = ", ".join(s["name"] for s in sheets)
    return {
        "ok": True,
        "kind": "xlsx",
        "path": str(file_path),
        "sheets": sheets,
        "message": (
            f"{len(sheets)} 个工作表: {names}。"
            "每个 sheet 含行/列数、投票表头、合并区域、公式别名与 markdown 预览。"
            "大 sheet（>5000 行）请配合 chunk_large_workbook 分块后再读。"
        ),
    }


# --- sheet_to_markdown ---------------------------------------------------------


def sheet_to_markdown(
    path: str,
    sheet: str = "",
    start_row: int = 1,
    end_row: int = 0,
    max_rows: int = _SLICE_MAX_ROWS,
) -> dict:
    """Export header plus an inclusive row range as a markdown table.

    ``end_row=0`` means the last row of the sheet. Rows outside the range
    are not included. Column width is the sheet's real used columns.
    """
    file_path, kind, err = _resolve_spreadsheet(path)
    if err is not None or file_path is None:
        return _denied(err, sheet=sheet or "", start_row=start_row, end_row=end_row)
    if start_row < 1:
        return {
            **_denied({"ok": False, "issue": "invalid_range", "message": "start_row 必须 >= 1 / start_row must be >= 1"}),
            "sheet": sheet or "",
            "start_row": start_row,
            "end_row": end_row,
        }
    if end_row < 0:
        return {
            **_denied({"ok": False, "issue": "invalid_range", "message": "end_row 必须 >= 0 / end_row must be >= 0"}),
            "sheet": sheet or "",
            "start_row": start_row,
            "end_row": end_row,
        }
    if end_row > 0 and end_row < start_row:
        return {
            **_denied({"ok": False, "issue": "invalid_range", "message": "行区间无效 / invalid row range"}),
            "sheet": sheet or "",
            "start_row": start_row,
            "end_row": end_row,
        }
    if max_rows < 1 or max_rows > _SLICE_MAX_ROWS:
        return {
            **_denied({
                "ok": False,
                "issue": "invalid_range",
                "message": (
                    f"max_rows 必须在 1..{_SLICE_MAX_ROWS} / max_rows must be within 1..{_SLICE_MAX_ROWS}"
                ),
            }),
            "sheet": sheet or "",
            "start_row": start_row,
            "end_row": end_row,
        }

    if kind == "csv":
        return _markdown_csv(
            file_path, start_row=start_row, end_row=end_row, max_rows=max_rows
        )
    return _markdown_xlsx(
        file_path,
        sheet=sheet,
        start_row=start_row,
        end_row=end_row,
        max_rows=max_rows,
    )


def _markdown_csv(
    file_path: Path, *, start_row: int, end_row: int, max_rows: int
) -> dict[str, Any]:
    from .guards import detect_encoding

    enc = detect_encoding(str(file_path))
    encoding = enc.get("encoding") or "utf-8"
    read_encoding = "utf-8" if encoding in {"", "unknown"} else encoding
    total = _count_text_lines(file_path)
    last = total if end_row == 0 else end_row
    if last < start_row:
        return {
            **_denied({"ok": False, "issue": "invalid_range", "message": "行区间无效 / invalid row range"}),
            "sheet": file_path.stem,
            "start_row": start_row,
            "end_row": end_row,
        }
    if end_row > 0 and last - start_row + 1 > _SLICE_MAX_ROWS:
        return {
            **_denied({
                "ok": False,
                "issue": "invalid_range",
                "message": (
                    f"行区间过大 / Row range too large: {last - start_row + 1} > "
                    f"{_SLICE_MAX_ROWS}。请用 chunk_large_workbook 分块后逐段读取。"
                ),
            }),
            "sheet": file_path.stem,
            "start_row": start_row,
            "end_row": end_row,
        }
    last = min(last, start_row + max_rows - 1, start_row + _SLICE_MAX_ROWS - 1)

    header: list[str] = []
    ranged: list[list[str]] = []
    try:
        with open(
            file_path, "r", encoding=read_encoding, errors="replace", newline=""
        ) as fh:
            for i, row in enumerate(csv.reader(fh), start=1):
                cells = [str(c) for c in row]
                if i == 1:
                    header = cells
                if i > last:
                    break
                if i >= start_row:
                    ranged.append(cells)
    except OSError as exc:
        return {
            "ok": False,
            "issue": "unreadable",
            "sheet": file_path.stem,
            "start_row": start_row,
            "end_row": end_row,
            "markdown": "",
            "message": f"无法读取文件 / Cannot read file: {exc}",
        }
    md_rows = ranged if start_row <= 1 else [header] + ranged
    return {
        "ok": True,
        "issue": "",
        "sheet": file_path.stem,
        "start_row": start_row,
        "end_row": last,
        "n_rows": len(ranged),
        "header_included": start_row > 1,
        "markdown": _rows_to_markdown(md_rows) if md_rows else "",
        "message": (
            f"已导出 Markdown / Markdown export: {_safe_name(file_path)} "
            f"rows={start_row}-{last}"
        ),
    }


def _markdown_xlsx(
    file_path: Path, *, sheet: str, start_row: int, end_row: int, max_rows: int
) -> dict[str, Any]:
    try:
        wb_v = load_workbook(file_path, read_only=True, data_only=True)
        wb_f = load_workbook(file_path, read_only=True, data_only=False)
    except Exception as exc:  # noqa: BLE001
        return {
            "ok": False,
            "issue": "unreadable",
            "sheet": sheet,
            "start_row": start_row,
            "end_row": end_row,
            "markdown": "",
            "message": (
                f"无法打开工作簿 / Cannot open workbook: {type(exc).__name__}: {exc}。"
                "请先调用 detect_corrupt_workbook。"
            ),
        }
    try:
        if sheet:
            if sheet not in wb_v.sheetnames:
                return {
                    "ok": False,
                    "issue": "missing_sheet",
                    "sheet": sheet,
                    "start_row": start_row,
                    "end_row": end_row,
                    "markdown": "",
                    "message": (
                        f"找不到工作表 / Sheet not found: {sheet}。"
                        f"可用: {', '.join(wb_v.sheetnames[:20])}"
                    ),
                }
            ws = wb_v[sheet]
        else:
            ws = wb_v[wb_v.sheetnames[0]]
        ws_f = wb_f[ws.title]
        last_row = int(ws.max_row or 0)
        last = last_row if end_row == 0 else end_row
        if last < start_row:
            return {
                "ok": False,
                "issue": "invalid_range",
                "sheet": ws.title,
                "start_row": start_row,
                "end_row": end_row,
                "markdown": "",
                "message": "行区间无效 / invalid row range",
            }
        if end_row > 0 and last - start_row + 1 > _SLICE_MAX_ROWS:
            return {
                "ok": False,
                "issue": "invalid_range",
                "sheet": ws.title,
                "start_row": start_row,
                "end_row": end_row,
                "markdown": "",
                "message": (
                    f"行区间过大 / Row range too large: {last - start_row + 1} > "
                    f"{_SLICE_MAX_ROWS}。请用 chunk_large_workbook 分块后逐段读取。"
                ),
            }
        last = min(last, start_row + max_rows - 1, start_row + _SLICE_MAX_ROWS - 1)
        scan_kwargs = _iter_kwargs(ws, values_only=True)
        header = [
            _cell_str(c)
            for c in _trim_trailing_empty(
                list(next(ws.iter_rows(min_row=1, max_row=1, **scan_kwargs), ()))
            )
        ]
        formulas = _formula_map_range(ws_f, start_row, last)
        ranged: list[list[str]] = []
        uncached = False
        for row_idx, row in enumerate(
            ws.iter_rows(min_row=start_row, max_row=last, **scan_kwargs),
            start=start_row,
        ):
            values = list(row)
            for c_idx in range(len(values)):
                if values[c_idx] is None:
                    ref = _cell_ref(row_idx, c_idx + 1)
                    if ref in formulas:
                        values[c_idx] = formulas[ref]
                        uncached = True
            ranged.append([_cell_str(c) for c in _trim_trailing_empty(values)])
        md_rows = ranged if start_row <= 1 else [header] + ranged
        return {
            "ok": True,
            "issue": "",
            "sheet": ws.title,
            "start_row": start_row,
            "end_row": last,
            "n_rows": len(ranged),
            "header_included": start_row > 1,
            "markdown": _rows_to_markdown(md_rows) if md_rows else "",
            "note": _UNCACHED_PREVIEW_NOTE if uncached else "",
            "message": (
                f"已导出 Markdown / Markdown export: {_safe_name(file_path)} "
                f"sheet={ws.title} rows={start_row}-{last}"
            ),
        }
    finally:
        wb_v.close()
        wb_f.close()


def _formula_map_range(ws: Any, start_row: int, end_row: int) -> dict[str, str]:
    """cell ref -> formula within one row range (slice-local, cheap)."""
    out: dict[str, str] = {}
    try:
        for row_idx, row in enumerate(
            ws.iter_rows(
                min_row=start_row,
                max_row=min(end_row, start_row + _FORMULA_SCAN_ROWS - 1),
                **_iter_kwargs(ws, values_only=False),
            ),
            start=start_row,
        ):
            for col_idx, cell in enumerate(row, start=1):
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    out[_cell_ref(row_idx, col_idx)] = value
    except Exception as exc:  # noqa: BLE001 — formulas are best-effort context
        logger.warning("formula scan failed: %s", exc)
    return out


# --- audit_workbook ------------------------------------------------------------


def audit_workbook(path: str, max_rows: int = _AUDIT_MAX_ROWS) -> dict:
    """Post-write self-check: classified formula errors and suspicious constants.

    Mirrors the "needs review" diff of the Shortcut writeup: cached error
    values (#REF! & co) are must-fix; numeric literals of 3+ digits embedded
    in formulas are flagged for review. Streaming, read-only, capped.
    """
    file_path, kind, err = _resolve_spreadsheet(path)
    if err is not None or file_path is None:
        return {"ok": False, "message": err["message"]}

    if kind == "csv":
        return {
            "ok": True,
            "kind": "csv",
            "n_formula_cells": 0,
            "must_fix": [],
            "review": [],
            "message": "CSV 无公式，跳过公式审计 / CSV has no formulas; audit skipped.",
        }

    if max_rows < 1:
        max_rows = 1
    max_rows = min(max_rows, _AUDIT_MAX_ROWS)

    try:
        wb_v = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return {
            "ok": False,
            "message": (
                f"无法打开工作簿 / Cannot open workbook: {type(exc).__name__}: {exc}。"
                "请先调用 detect_corrupt_workbook。"
            ),
        }
    wb_f = load_workbook(file_path, read_only=True, data_only=False)

    must_fix: list[dict[str, str]] = []
    review: list[dict[str, str]] = []
    n_formula_cells = 0
    n_cells = 0
    truncated = False
    try:
        for name in wb_v.sheetnames[:_MD_MAX_SHEETS]:
            ws_v = wb_v[name]
            ws_f = wb_f[name]
            f_iter = ws_f.iter_rows(max_row=max_rows, max_col=_MD_MAX_COLS)
            for row_idx, row_v in enumerate(
                ws_v.iter_rows(max_row=max_rows, max_col=_MD_MAX_COLS, values_only=True),
                start=1,
            ):
                try:
                    row_f = next(f_iter)
                except StopIteration:
                    row_f = []
                for col_idx in range(max(len(row_v), len(row_f))):
                    v = row_v[col_idx] if col_idx < len(row_v) else None
                    f = row_f[col_idx].value if col_idx < len(row_f) else None
                    if isinstance(f, str) and f.startswith("="):
                        n_formula_cells += 1
                    n_cells += 1
                    ref = _cell_ref(row_idx, col_idx + 1)
                    if isinstance(v, str) and v in _EXCEL_ERRORS:
                        must_fix.append(
                            {
                                "sheet": name,
                                "cell": ref,
                                "error": v,
                                "formula": f if isinstance(f, str) else "",
                            }
                        )
                        if len(must_fix) >= _AUDIT_MAX_ENTRIES:
                            truncated = True
                    if (
                        isinstance(f, str)
                        and f.startswith("=")
                        and _FORMULA_LITERAL_RE.search(f)
                        and len(review) < _AUDIT_MAX_ENTRIES
                    ):
                        review.append({"sheet": name, "cell": ref, "formula": f})
                if row_idx >= max_rows:
                    truncated = truncated or int(ws_v.max_row or 0) > max_rows
                    break
    finally:
        wb_v.close()
        wb_f.close()

    must_fix = must_fix[:_AUDIT_MAX_ENTRIES]
    review = review[:_AUDIT_MAX_ENTRIES]
    parts = [f"扫描 {n_cells} 个单元格，其中公式 {n_formula_cells} 个。"]
    if must_fix:
        parts.append(
            f"必须修复 {len(must_fix)} 处公式错误（如 #REF!）: "
            + "; ".join(f"{e['sheet']}!{e['cell']}={e['error']}" for e in must_fix[:5])
            + ("…" if len(must_fix) > 5 else "")
        )
    else:
        parts.append("未发现公式错误值。")
    if review:
        parts.append(
            f"疑似硬编码待复核 {len(review)} 处: "
            + "; ".join(f"{e['sheet']}!{e['cell']} {e['formula'][:40]}" for e in review[:3])
            + ("…" if len(review) > 3 else "")
        )
    if truncated:
        parts.append(f"已达扫描上限 {max_rows} 行，结果可能不完整。")
    return {
        "ok": True,
        "kind": "xlsx",
        "n_cells": n_cells,
        "n_formula_cells": n_formula_cells,
        "must_fix": must_fix,
        "review": review,
        "truncated": truncated,
        "message": " ".join(parts),
    }


# --- preflight_workbook --------------------------------------------------------


def preflight_workbook(path: str, max_rows: int = 5000) -> dict:
    """One-call guard: corrupt check + encoding + chunk plan, short-circuit on corruption.

    Collapses the three fixed guard calls every task was paying into one
    round trip. Prefer this over calling the three guard tools separately;
    fall back to the individual tools only when you need to re-check after
    an edit. Structure/preview context still comes from ``describe_workbook``.
    """
    from .guards import chunk_large_workbook, detect_corrupt_workbook, detect_encoding

    corrupt = detect_corrupt_workbook(path)
    if not corrupt.get("ok"):
        return {
            "ok": False,
            "stage": "corrupt",
            "proceed": False,
            "corrupt": corrupt,
            "encoding": None,
            "chunking": None,
            "message": corrupt.get("message") or "文件损坏，终止问答。",
        }

    encoding = detect_encoding(path)
    chunking = chunk_large_workbook(path, max_rows=max_rows)
    encoding_hint = (
        encoding.get("encoding")
        if encoding.get("encoding") not in (None, "unknown")
        else "utf-8"
    )
    proceed = True
    notes: list[str] = []
    if float(encoding.get("confidence") or 0.0) < 0.5 and encoding.get("encoding") not in (
        "utf-8", "ascii"
    ):
        notes.append("编码置信度偏低，读取前请先转码为 UTF-8。")
    if chunking.get("needs_chunking"):
        notes.append(
            f"最大 sheet {chunking.get('max_sheet_rows')} 行，"
            f"请按 {len(chunking.get('chunks') or [])} 个分块区间逐段处理。"
        )
    return {
        "ok": True,
        "stage": "done",
        "proceed": proceed,
        "corrupt": corrupt,
        "encoding": {"encoding": encoding_hint, "confidence": encoding.get("confidence")},
        "chunking": chunking,
        "message": (
            "护栏通过："
            f"encoding={encoding_hint}"
            + ("；" + " ".join(notes) if notes else "，无需分块。")
        ),
    }
