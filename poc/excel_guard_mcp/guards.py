# -*- coding: utf-8 -*-
"""Excel anomaly guards for POC demo: corrupt / encoding / large-sheet chunking."""

from __future__ import annotations

import os
import zipfile
from pathlib import Path

import chardet
from openpyxl import load_workbook

# Max bytes fed to chardet (avoid OOM on huge CSV before chunking).
_ENCODING_SAMPLE_BYTES = 1 << 20  # 1 MiB
_ENCODING_HARD_LIMIT = 32 << 20  # 32 MiB — refuse full detect beyond this


def _workspace_root() -> Path:
    raw = os.environ.get("POC_WORKSPACE") or os.environ.get("QWENPAW_WORKING_DIR")
    if raw:
        return Path(raw).expanduser().resolve()
    # Default: repository root (parent of poc/)
    return Path(__file__).resolve().parents[2]


def _resolve_allowed_path(path: str) -> tuple[Path | None, dict | None]:
    """Resolve path and enforce workspace sandbox.

    Returns (path, None) on success, or (None, error_dict) on denial/missing shape.
    error_dict uses ok/issue/message keys for corrupt-style callers; encoding/chunk
    callers remap as needed.
    """
    root = _workspace_root()
    try:
        file_path = Path(path).expanduser().resolve(strict=False)
    except OSError as exc:
        return None, {
            "ok": False,
            "issue": "path_denied",
            "message": f"路径无法解析 / Cannot resolve path: {exc}",
        }

    try:
        if not file_path.is_relative_to(root):
            return None, {
                "ok": False,
                "issue": "path_denied",
                "message": (
                    f"路径越界 / Path outside workspace root ({root}): {path}"
                ),
            }
    except AttributeError:
        # Python < 3.9 fallback — not expected on 3.13+
        if root not in file_path.parents and file_path != root:
            return None, {
                "ok": False,
                "issue": "path_denied",
                "message": (
                    f"路径越界 / Path outside workspace root ({root}): {path}"
                ),
            }

    if file_path.is_symlink():
        try:
            real = file_path.resolve(strict=True)
            if not real.is_relative_to(root):
                return None, {
                    "ok": False,
                    "issue": "path_denied",
                    "message": (
                        f"符号链接越界 / Symlink escapes workspace: {path}"
                    ),
                }
            file_path = real
        except OSError as exc:
            return None, {
                "ok": False,
                "issue": "path_denied",
                "message": f"符号链接无效 / Invalid symlink: {exc}",
            }

    return file_path, None


def detect_corrupt_workbook(path: str) -> dict:
    """Detect whether an Excel workbook is corrupt or unreadable.

    Returns:
        {"ok": bool, "issue": str, "message": str}
    """
    file_path, err = _resolve_allowed_path(path)
    if err is not None:
        return err

    assert file_path is not None
    if not file_path.exists():
        return {
            "ok": False,
            "issue": "missing",
            "message": (
                f"文件不存在 / File not found: {path}。"
                "请确认路径是否正确。"
            ),
        }

    suffix = file_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        try:
            if not zipfile.is_zipfile(file_path):
                return {
                    "ok": False,
                    "issue": "corrupt",
                    "message": (
                        "检测到损坏的 Excel 工作簿 / Corrupt workbook detected: "
                        f"{file_path.name}。文件不是有效的 Office Open XML (ZIP) 包，"
                        "无法用 openpyxl 打开。请重新导出或修复后再试。"
                    ),
                }
            with zipfile.ZipFile(file_path, "r") as zf:
                names = zf.namelist()
                if "[Content_Types].xml" not in names:
                    return {
                        "ok": False,
                        "issue": "corrupt",
                        "message": (
                            "检测到损坏的 Excel 工作簿 / Corrupt workbook detected: "
                            f"{file_path.name}。ZIP 包缺少 [Content_Types].xml，"
                            "结构不完整。请重新导出或修复后再试。"
                        ),
                    }
            load_workbook(file_path, read_only=True, data_only=True).close()
        except Exception as exc:  # noqa: BLE001 - surface open failure as corrupt
            return {
                "ok": False,
                "issue": "corrupt",
                "message": (
                    "检测到损坏的 Excel 工作簿 / Corrupt workbook detected: "
                    f"{file_path.name}。打开失败：{type(exc).__name__}。"
                    "请重新导出或修复后再试。"
                ),
            }
        return {
            "ok": True,
            "issue": "",
            "message": f"工作簿正常 / Workbook OK: {file_path.name}",
        }

    try:
        load_workbook(file_path, read_only=True, data_only=True).close()
    except Exception as exc:  # noqa: BLE001
        return {
            "ok": False,
            "issue": "corrupt",
            "message": (
                "检测到损坏或不支持的 Excel 文件 / Corrupt or unsupported workbook: "
                f"{file_path.name}。打开失败：{type(exc).__name__}。"
                "请使用有效的 .xlsx 文件。"
            ),
        }
    return {
        "ok": True,
        "issue": "",
        "message": f"工作簿正常 / Workbook OK: {file_path.name}",
    }


def detect_encoding(path: str) -> dict:
    """Detect text/CSV encoding with chardet; describe xlsx as zip/utf-8 internals.

    Returns:
        {"encoding": str, "confidence": float, "message": str}
    """
    file_path, err = _resolve_allowed_path(path)
    if err is not None:
        return {
            "encoding": "unknown",
            "confidence": 0.0,
            "message": err["message"],
        }

    assert file_path is not None
    if not file_path.exists():
        return {
            "encoding": "unknown",
            "confidence": 0.0,
            "message": f"文件不存在 / File not found: {path}",
        }

    suffix = file_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        is_zip = zipfile.is_zipfile(file_path)
        return {
            "encoding": "utf-8",
            "confidence": 1.0 if is_zip else 0.0,
            "message": (
                f"XLSX 为 ZIP 包，内部 XML 通常为 UTF-8 / "
                f"XLSX is a ZIP package with UTF-8 XML internals: {file_path.name}。"
                + (
                    "可直接用 openpyxl/pandas 读取。"
                    if is_zip
                    else "但当前文件不是有效 ZIP，请先做损坏检测。"
                )
            ),
        }

    size = file_path.stat().st_size
    if size == 0:
        return {
            "encoding": "utf-8",
            "confidence": 0.0,
            "message": f"空文件 / Empty file: {file_path.name}",
        }
    if size > _ENCODING_HARD_LIMIT:
        return {
            "encoding": "unknown",
            "confidence": 0.0,
            "message": (
                f"文件过大无法做全量编码探测 / File too large for encoding detect: "
                f"{file_path.name} ({size} bytes > {_ENCODING_HARD_LIMIT})。"
                "请先转码为 UTF-8 或拆分后再测。"
            ),
        }

    with file_path.open("rb") as fh:
        raw = fh.read(_ENCODING_SAMPLE_BYTES)

    detected = chardet.detect(raw) or {}
    encoding = (detected.get("encoding") or "unknown").lower()
    confidence = float(detected.get("confidence") or 0.0)

    if encoding in {"ascii", "utf-8", "utf8"}:
        try:
            raw.decode("utf-8")
            return {
                "encoding": "utf-8" if encoding != "ascii" else "ascii",
                "confidence": confidence,
                "message": (
                    f"文本编码正常 / Encoding OK ({encoding}, confidence={confidence:.2f}): "
                    f"{file_path.name}"
                ),
            }
        except UnicodeDecodeError:
            pass

    if confidence < 0.5:
        tip = (
            f"编码置信度偏低 / Low-confidence encoding guess: {encoding} "
            f"(confidence={confidence:.2f})，文件：{file_path.name}。"
            "建议按 latin-1/gbk 尝试转码为 UTF-8 后再读取，勿直接当 UTF-8 打开。"
        )
    else:
        tip = (
            f"检测到非 UTF-8 编码 / Non-UTF-8 encoding detected: {encoding} "
            f"(confidence={confidence:.2f})，文件：{file_path.name}。"
            "读取 CSV/文本前请先转码为 UTF-8，否则中文或特殊字符可能乱码。"
        )
    return {
        "encoding": encoding,
        "confidence": confidence,
        "message": tip,
    }


def chunk_large_workbook(path: str, max_rows: int = 5000) -> dict:
    """Plan row-range chunks when a workbook sheet exceeds max_rows.

    Returns:
        {
            "needs_chunking": bool,
            "total_rows": int,
            "chunks": list[{"start_row": int, "end_row": int, "sheet": str}],
            "message": str,
        }
    """
    file_path, err = _resolve_allowed_path(path)
    if err is not None:
        return {
            "needs_chunking": False,
            "total_rows": 0,
            "chunks": [],
            "message": err["message"],
        }

    assert file_path is not None
    if not file_path.exists():
        return {
            "needs_chunking": False,
            "total_rows": 0,
            "chunks": [],
            "message": f"文件不存在 / File not found: {path}",
        }

    if max_rows < 1:
        return {
            "needs_chunking": False,
            "total_rows": 0,
            "chunks": [],
            "message": "max_rows 必须 >= 1 / max_rows must be >= 1",
        }

    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return {
            "needs_chunking": False,
            "total_rows": 0,
            "chunks": [],
            "message": (
                "无法打开工作簿做分块 / Cannot open workbook for chunking: "
                f"{type(exc).__name__}"
            ),
        }

    chunks: list[dict] = []
    total_rows = 0
    try:
        for sheet in wb.worksheets:
            sheet_rows = int(sheet.max_row or 0)
            total_rows = max(total_rows, sheet_rows)
            if sheet_rows <= max_rows:
                continue
            start = 1
            while start <= sheet_rows:
                end = min(start + max_rows - 1, sheet_rows)
                chunks.append(
                    {
                        "start_row": start,
                        "end_row": end,
                        "sheet": sheet.title,
                    }
                )
                start = end + 1
    finally:
        wb.close()

    needs_chunking = total_rows > max_rows
    if needs_chunking:
        message = (
            f"工作表行数过大，需要分块处理 / Large sheet needs chunking: "
            f"total_rows={total_rows} > max_rows={max_rows}，"
            f"已生成 {len(chunks)} 个分块范围。"
        )
    else:
        message = (
            f"行数未超限，无需分块 / No chunking needed: "
            f"total_rows={total_rows} <= max_rows={max_rows}"
        )
        chunks = []

    return {
        "needs_chunking": needs_chunking,
        "total_rows": total_rows,
        "chunks": chunks,
        "message": message,
    }
