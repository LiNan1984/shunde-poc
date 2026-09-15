# -*- coding: utf-8 -*-
"""Real tests for poc.report_mcp (no mocks of subject logic)."""

from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image  # matplotlib ships with Pillow

from poc.report_mcp import (
    pivot_table,
    render_bar,
    render_docx_report,
    render_heatmap,
    render_line,
    render_pie,
    render_scatter,
)
from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(autouse=True)
def _sandbox_workspace(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Confine all path checks to the per-test temp dir."""
    monkeypatch.setenv("POC_WORKSPACE", str(tmp_path))


def test_pivot_table_happy_path() -> None:
    data = [
        {"region": "N", "q": 1, "sales": 10},
        {"region": "S", "q": 1, "sales": 20},
        {"region": "N", "q": 2, "sales": 12},
        {"region": "S", "q": 2, "sales": 22},
    ]
    res = pivot_table(data, rows=["q"], cols=["region"], value="sales", aggfunc="sum")
    assert res["ok"] is True
    assert res["row_count"] == 2
    assert res["col_count"] == 2
    assert isinstance(res["table"], list)
    assert len(res["table"]) >= 2


def test_pivot_table_rejects_too_many_categories() -> None:
    data = [{"x": i, "y": i % 5, "v": 1} for i in range(200)]
    res = pivot_table(data, rows=["x"], cols=["y"], value="v", aggfunc="count", max_categories=10)
    assert res["ok"] is False
    assert "超限" in res["message"] or "too large" in res["message"].lower()


def test_render_bar_produces_readable_png(tmp_path: Path) -> None:
    out = tmp_path / "charts" / "bar.png"
    data = [{"region": r, "sales": v} for r, v in zip("NSEW", [10, 20, 15, 5])]
    res = render_bar(data, x="region", y="sales", title="Sales", output_path=str(out))
    assert res["ok"] is True, res
    assert Path(res["output_path"]).is_file()
    with Image.open(res["output_path"]) as img:
        assert img.format == "PNG"
        assert img.width > 50 and img.height > 50


def test_render_heatmap_from_matrix(tmp_path: Path) -> None:
    out = tmp_path / "charts" / "hm.png"
    res = render_heatmap(
        matrix=[[1, 2, 3], [4, 5, 6]],
        x_labels=["A", "B", "C"],
        y_labels=["r1", "r2"],
        output_path=str(out),
    )
    assert res["ok"] is True, res
    assert Path(res["output_path"]).is_file()


def test_render_line_pie_scatter_smoke(tmp_path: Path) -> None:
    data = [{"x": i, "y": i * 2, "k": "a" if i < 2 else "b"} for i in range(4)]
    r_line = render_line(data, x="x", y="y", output_path=str(tmp_path / "line.png"))
    r_pie = render_pie(data, label="k", value="y", output_path=str(tmp_path / "pie.png"))
    r_scatter = render_scatter(data, x="x", y="y", output_path=str(tmp_path / "sc.png"))
    for r in (r_line, r_pie, r_scatter):
        assert r["ok"] is True, r
        assert Path(r["output_path"]).is_file()


def test_render_docx_report_round_trip(tmp_path: Path) -> None:
    chart = tmp_path / "chart.png"
    chart.parent.mkdir(parents=True, exist_ok=True)
    bar = render_bar(
        [{"k": "A", "v": 1}, {"k": "B", "v": 3}],
        x="k",
        y="v",
        output_path=str(chart),
    )
    assert bar["ok"] is True

    docx_path = tmp_path / "report.docx"
    sections = [
        {
            "heading": "概述",
            "level": 2,
            "paragraphs": ["这是第一节正文。This is the first section body."],
            "table": {"headers": ["k", "v"], "rows": [["A", 1], ["B", 3]]},
            "charts": [bar["output_path"]],
        }
    ]
    res = render_docx_report(
        title="分行销售报告",
        sections=sections,
        output_path=str(docx_path),
        header="顺德农商行 POC",
        footer="演示页",
    )
    assert res["ok"] is True, res
    assert Path(res["output_path"]).is_file()

    # Round-trip: reopen the docx and check structural contents.
    doc = Document(str(docx_path))
    texts = "\n".join(p.text for p in doc.paragraphs)
    assert "分行销售报告" in texts
    assert "概述" in texts
    assert "这是第一节正文" in texts
    # Tables: at least one inserted.
    assert len(doc.tables) >= 1
    # Inline images: at least one shape with an inline drawing element.
    from docx.oxml.ns import qn

    drawings = doc.element.findall(".//" + qn("w:drawing"))
    assert len(drawings) >= 1


def test_path_outside_workspace_denied(tmp_path: Path) -> None:
    outside = Path("/etc/evil_report.docx")
    res = render_docx_report(
        title="x",
        sections=[],
        output_path=str(outside),
    )
    assert res["ok"] is False
    assert "越界" in res["message"] or "outside" in res["message"].lower()


def test_cjk_font_configured() -> None:
    import matplotlib.pyplot as plt

    from poc.report_mcp import charts

    assert charts._CJK_FONT, "no CJK font resolved on this machine"
    assert charts._CJK_FONT in plt.rcParams["font.sans-serif"]
    assert plt.rcParams["axes.unicode_minus"] is False


def test_render_bar_chinese_no_missing_glyph(tmp_path: Path, recwarn) -> None:
    data = [{"分行": "顺德支行", "金额": 120.5}, {"分行": "大良支行", "金额": 98.3}]
    res = render_bar(
        data,
        x="分行",
        y="金额",
        title="各分行金额（万元）",
        output_path=str(tmp_path / "cjk_bar.png"),
    )
    assert res["ok"] is True, res
    glyph_warnings = [w for w in recwarn.list if "missing from font" in str(w.message)]
    assert not glyph_warnings


def test_server_module_imports() -> None:
    from poc.report_mcp.server import mcp

    assert mcp is not None
