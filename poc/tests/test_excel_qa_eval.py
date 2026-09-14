# -*- coding: utf-8 -*-
"""Golden QA set invariants + scorer behavior (no mocks; files are real)."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from poc.evals.qa_set import build_qa_fixture, golden_answers, golden_questions
from poc.evals.score import score_answers


@pytest.fixture(autouse=True)
def _sandbox(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("POC_WORKSPACE", str(tmp_path))
    monkeypatch.setenv("POC_LOAD_SECRETS", "0")


def test_build_qa_fixture_writes_files_and_manifest(tmp_path: Path) -> None:
    built = build_qa_fixture(tmp_path)

    assert built["version"] == "1.0"
    assert len(built["questions"]) == 10
    for name in built["files"]:
        assert (tmp_path / name).is_file()
    manifest = json.loads((tmp_path / "qa_golden.json").read_text(encoding="utf-8"))
    assert manifest["version"] == "1.0"
    assert {q["id"] for q in manifest["questions"]} == {
        q["id"] for q in golden_questions()
    }


def test_golden_answers_recomputable_from_written_files(tmp_path: Path) -> None:
    """Every expected value must be derivable from the files on disk via a
    separate read path (pandas/openpyxl), not just from the in-memory lists."""
    build_qa_fixture(tmp_path)
    import pandas as pd
    from openpyxl import load_workbook

    golden = golden_answers()

    sales = pd.read_excel(tmp_path / "销售明细.xlsx", sheet_name="销售明细")
    assert round(float(sales[sales["地区"] == "顺德"]["销售额"].sum()), 2) == golden["q01_shunde_total"]
    assert len(sales) == golden["q02_row_count"]
    assert sales.loc[sales["销售额"].idxmax(), "网点"] == golden["q03_top_branch"]
    by_region = sales.groupby("地区")["销售额"].sum()
    assert by_region.idxmax() == golden["q04_top_region"]
    gz_jan = sales[(sales["地区"] == "广州") & sales["日期"].astype(str).str.startswith("2024-01")]
    assert round(float(gz_jan["销售额"].sum()), 2) == golden["q05_gz_jan_total"]

    wb = load_workbook(tmp_path / "分区域汇总.xlsx", data_only=True)
    ws = wb["分区域汇总"]
    rows = {row[0]: row for row in ws.iter_rows(min_row=4, values_only=True) if row[0]}
    assert rows["顺德"][1] == golden["q06_sd_2024_sales"]
    assert rows["广州"][4] == golden["q07_gz_2025_orders"]

    products = pd.read_excel(tmp_path / "公式列.xlsx", sheet_name="收入计算")
    revenue = products["数量"] * products["单价"]
    assert round(float(revenue.sum()), 2) == golden["q08_revenue_total"]
    assert int((revenue > 5000).sum()) == golden["q09_revenue_over_5000"]

    staff = pd.read_csv(tmp_path / "网点名单.csv", encoding="gbk")
    assert int((staff["部门"] == "零售部").sum()) == golden["q10_retail_headcount"]


def test_score_answers_perfect_golden_scores_one() -> None:
    questions = golden_questions()

    report = score_answers(questions, golden_answers())

    assert report["accuracy"] == 1.0
    assert report["correct"] == report["n"] == 10
    assert report["failed"] == []


def test_score_answers_flags_missing_and_wrong() -> None:
    questions = golden_questions()
    answers = golden_answers()
    answers["q02_row_count"] = 999
    answers["q03_top_branch"] = "不存在的网点"
    answers["q10_retail_headcount"] = "零售部共 12 人"  # string-wrapped number still passes
    del answers["q05_gz_jan_total"]

    report = score_answers(questions, answers)

    assert report["accuracy"] < 1.0
    failed_ids = {f["id"] for f in report["failed"]}
    assert failed_ids == {"q02_row_count", "q03_top_branch", "q05_gz_jan_total"}
    missing = next(f for f in report["failed"] if f["id"] == "q05_gz_jan_total")
    assert missing["got"] is None


def test_score_answers_numeric_tolerates_noise() -> None:
    questions = golden_questions()
    answers = golden_answers()
    answers["q01_shunde_total"] = f"{golden_answers()['q01_shunde_total']:,.2f} 元"

    report = score_answers(questions, answers)

    assert report["accuracy"] == 1.0


def test_preflight_shortcircuits_on_corrupt(tmp_path: Path) -> None:
    from poc.excel_guard_mcp.readers import preflight_workbook

    bad = tmp_path / "corrupt.xlsx"
    bad.write_bytes(b"NOT_A_VALID_XLSX")

    result = preflight_workbook(str(bad))

    assert result["ok"] is False
    assert result["stage"] == "corrupt"
    assert result["proceed"] is False
    assert result["encoding"] is None


def test_preflight_single_call_covers_three_guards(tmp_path: Path) -> None:
    from poc.excel_guard_mcp.readers import preflight_workbook
    from poc.fixtures.generate_fixtures import write_encoding_latin1_csv

    csv_path = tmp_path / "latin1.csv"
    write_encoding_latin1_csv(csv_path)

    result = preflight_workbook(str(csv_path))

    assert result["ok"] is True
    assert result["proceed"] is True
    assert result["corrupt"]["ok"] is True
    assert result["encoding"]["encoding"].lower().startswith("latin") or \
        result["encoding"]["encoding"] == "windows-1252"
    assert result["chunking"]["needs_chunking"] is False
