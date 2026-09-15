# -*- coding: utf-8 -*-
"""Integration tests for the FastMCP ``excel-guard`` server.

These tests exercise the real FastMCP tool registry and invoke the tools
through ``mcp.call_tool`` (no stdio transport is ever started). Guard logic
is not mocked; real xlsx/csv files are created under ``tmp_path`` and the
workspace sandbox is pinned to that directory via ``POC_WORKSPACE``.

The FastMCP API is coroutine-based; ``asyncio.run`` drives each call inside
plain synchronous pytest functions so no async pytest plugin is required.
"""

from __future__ import annotations

import asyncio
import json
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from poc.excel_guard_mcp.server import mcp, main

EXPECTED_TOOLS = {
    "detect_corrupt_workbook",
    "detect_encoding",
    "chunk_large_workbook",
    "describe_workbook",
    "sheet_to_markdown",
}
REPO_ROOT = Path(__file__).resolve().parents[2]


# ---------------------------------------------------------------------------
# Fixtures / helpers
# ---------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def _sandbox_workspace(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Confine every guard call to the per-test temp workspace."""
    monkeypatch.setenv("POC_WORKSPACE", str(tmp_path))


@pytest.fixture
def outside_workspace_file() -> str:
    """A real file that lives in a sibling temp dir, outside the sandbox."""
    fd, raw_path = tempfile.mkstemp(prefix="excel-guard-outside-", suffix=".csv")
    try:
        os.write(fd, b"id,name\n1,alice\n")
    finally:
        os.close(fd)
    yield raw_path
    Path(raw_path).unlink(missing_ok=True)


def _make_xlsx(path: Path, rows: int = 1, sheet_title: str = "Sheet1") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(["id", "value"])
    for i in range(max(0, rows - 1)):
        ws.append([i, f"v{i}"])
    wb.save(path)
    return path


def _list_tools() -> dict[str, Any]:
    """Return registered tools keyed by name."""
    tools = asyncio.run(mcp.list_tools())
    return {tool.name: tool for tool in tools}


def _call_tool(name: str, arguments: dict[str, Any]) -> dict[str, Any]:
    """Invoke a registered MCP tool and decode its JSON payload."""
    result = asyncio.run(mcp.call_tool(name, arguments))
    if isinstance(result, dict):  # forward-compatible structured return
        return result
    assert result, f"tool {name!r} returned no content blocks"
    text = getattr(result[0], "text", None)
    assert text is not None, f"tool {name!r} returned non-text block: {result[0]!r}"
    payload = json.loads(text)
    assert isinstance(payload, dict)
    return payload


# ---------------------------------------------------------------------------
# 1. Tool registration (unit level)
# ---------------------------------------------------------------------------


def test_server_registers_exactly_five_tools() -> None:
    tools = _list_tools()

    assert len(tools) == 5
    assert set(tools) == EXPECTED_TOOLS


def test_mcp_describe_markdown_are_readers_callables() -> None:
    from poc.excel_guard_mcp import readers, server

    assert server._describe_workbook is readers.describe_workbook
    assert server._sheet_to_markdown is readers.sheet_to_markdown
    assert "audit_workbook" not in _list_tools()


def test_server_name_is_excel_guard() -> None:
    assert mcp.name == "excel-guard"


@pytest.mark.parametrize(
    ("tool_name", "keyword"),
    [
        ("detect_corrupt_workbook", "损坏"),
        ("detect_encoding", "编码"),
        ("chunk_large_workbook", "分块"),
        ("describe_workbook", "sheet"),
        ("sheet_to_markdown", "markdown"),
    ],
)
def test_registered_tools_have_descriptions(tool_name: str, keyword: str) -> None:
    tool = _list_tools()[tool_name]

    assert tool.description
    assert tool.description.strip()
    assert keyword in tool.description.lower()


def test_tool_schemas_require_path() -> None:
    tools = _list_tools()

    for tool_name in EXPECTED_TOOLS:
        schema = tools[tool_name].inputSchema
        assert schema["type"] == "object"
        assert schema.get("required") == ["path"]
        assert schema["properties"]["path"]["type"] == "string"


def test_chunk_tool_schema_declares_max_rows_default() -> None:
    schema = _list_tools()["chunk_large_workbook"].inputSchema
    max_rows = schema["properties"]["max_rows"]

    assert max_rows["type"] == "integer"
    assert max_rows["default"] == 5000
    # max_rows is optional; only path is required
    assert "max_rows" not in schema["required"]


# ---------------------------------------------------------------------------
# 2. Tool execution through the MCP server (integration level)
# ---------------------------------------------------------------------------


def test_mcp_detect_corrupt_workbook_accepts_valid_xlsx(tmp_path: Path) -> None:
    good = _make_xlsx(tmp_path / "ok.xlsx")

    payload = _call_tool("detect_corrupt_workbook", {"path": str(good)})

    assert payload["ok"] is True
    assert payload["issue"] == ""
    assert "OK" in payload["message"] or "正常" in payload["message"]


def test_mcp_detect_corrupt_workbook_flags_corrupt_file(tmp_path: Path) -> None:
    bad = tmp_path / "broken.xlsx"
    bad.write_bytes(b"this is definitely not a zip / xlsx package")

    payload = _call_tool("detect_corrupt_workbook", {"path": str(bad)})

    assert payload["ok"] is False
    assert payload["issue"] == "corrupt"
    assert "损坏" in payload["message"] or "Corrupt" in payload["message"]


def test_mcp_detect_encoding_reports_csv_encoding(tmp_path: Path) -> None:
    csv_path = tmp_path / "latin1.csv"
    csv_path.write_bytes(b"name,city\nAlice,Caf\xe9\n")

    payload = _call_tool("detect_encoding", {"path": str(csv_path)})

    assert set(payload) >= {"encoding", "confidence", "message"}
    assert isinstance(payload["encoding"], str) and payload["encoding"]
    assert isinstance(payload["confidence"], (int, float))
    assert 0.0 <= float(payload["confidence"]) <= 1.0
    # Latin-1 sample must not be misreported as UTF-8
    assert payload["encoding"].lower().replace("_", "-") not in {"utf-8", "utf8"}
    assert "UTF-8" in payload["message"] or "utf-8" in payload["message"].lower()


def test_mcp_detect_encoding_describes_xlsx_as_zip_utf8(tmp_path: Path) -> None:
    xlsx = _make_xlsx(tmp_path / "note.xlsx")

    payload = _call_tool("detect_encoding", {"path": str(xlsx)})

    assert payload["encoding"].lower() in {"utf-8", "utf8"}
    assert payload["confidence"] == 1.0
    assert "ZIP" in payload["message"] or "zip" in payload["message"]


def test_mcp_chunk_large_workbook_plans_contiguous_chunks(tmp_path: Path) -> None:
    large = _make_xlsx(tmp_path / "large.xlsx", rows=121, sheet_title="Sheet1")

    payload = _call_tool(
        "chunk_large_workbook", {"path": str(large), "max_rows": 50}
    )

    assert payload["needs_chunking"] is True
    assert payload["total_rows"] == 121
    chunks = payload["chunks"]
    assert len(chunks) == 3
    assert [c["start_row"] for c in chunks] == [1, 51, 101]
    assert [c["end_row"] for c in chunks] == [50, 100, 121]
    # Ranges must be contiguous and cover the sheet exactly
    for prev, curr in zip(chunks, chunks[1:]):
        assert curr["start_row"] == prev["end_row"] + 1
    assert chunks[-1]["end_row"] == payload["total_rows"]
    assert all(c["sheet"] == "Sheet1" for c in chunks)
    assert "分块" in payload["message"] or "chunk" in payload["message"].lower()


def test_mcp_chunk_large_workbook_uses_default_max_rows(tmp_path: Path) -> None:
    small = _make_xlsx(tmp_path / "small.xlsx", rows=3)

    payload = _call_tool("chunk_large_workbook", {"path": str(small)})

    assert payload["needs_chunking"] is False
    assert payload["total_rows"] == 3
    assert payload["chunks"] == []
    assert "无需分块" in payload["message"] or "No chunking" in payload["message"]
    assert "max_rows=5000" in payload["message"]


@pytest.mark.parametrize(
    ("tool_name", "checks"),
    [
        (
            "detect_corrupt_workbook",
            lambda p: [
                p["ok"] is False,
                p["issue"] == "path_denied",
                "路径越界" in p["message"],
            ],
        ),
        (
            "detect_encoding",
            lambda p: [
                p["encoding"] == "unknown",
                p["confidence"] == 0.0,
                "路径越界" in p["message"],
            ],
        ),
        (
            "chunk_large_workbook",
            lambda p: [
                p["needs_chunking"] is False,
                p["total_rows"] == 0,
                p["chunks"] == [],
                "路径越界" in p["message"],
            ],
        ),
        (
            "describe_workbook",
            lambda p: [
                p["ok"] is False,
                p["issue"] == "path_denied",
                "路径越界" in p["message"],
            ],
        ),
        (
            "sheet_to_markdown",
            lambda p: [
                p["ok"] is False,
                p["issue"] == "path_denied",
                "路径越界" in p["message"],
            ],
        ),
    ],
)
def test_mcp_tools_deny_paths_outside_workspace(
    tool_name: str, checks, outside_workspace_file: str
) -> None:
    payload = _call_tool(tool_name, {"path": outside_workspace_file})

    for check in checks(payload):
        assert check
    if tool_name in {"describe_workbook", "sheet_to_markdown"}:
        assert Path(outside_workspace_file).name not in payload["message"]


def test_mcp_detect_corrupt_reports_missing_file(tmp_path: Path) -> None:
    missing = tmp_path / "does-not-exist.xlsx"

    payload = _call_tool("detect_corrupt_workbook", {"path": str(missing)})

    assert payload["ok"] is False
    assert payload["issue"] == "missing"
    assert "不存在" in payload["message"] or "not found" in payload["message"].lower()


def test_mcp_detect_corrupt_zip_without_content_types(tmp_path: Path) -> None:
    """A ZIP archive that is valid structurally but lacks the OOXML manifest."""
    fake_xlsx = tmp_path / "no-manifest.xlsx"
    with zipfile.ZipFile(fake_xlsx, "w") as zf:
        zf.writestr("hello.txt", "not an office package")

    payload = _call_tool("detect_corrupt_workbook", {"path": str(fake_xlsx)})

    assert payload["ok"] is False
    assert payload["issue"] == "corrupt"
    assert "[Content_Types].xml" in payload["message"]


def test_mcp_detect_corrupt_zip_with_bad_internals(tmp_path: Path) -> None:
    """ZIP + manifest present, but openpyxl still cannot open it."""
    fake_xlsx = tmp_path / "bad-internals.xlsx"
    with zipfile.ZipFile(fake_xlsx, "w") as zf:
        zf.writestr("[Content_Types].xml", "<not-really-content-types/>")

    payload = _call_tool("detect_corrupt_workbook", {"path": str(fake_xlsx)})

    assert payload["ok"] is False
    # Bad internals: caught as corrupt or unreadable depending on exact failure path
    assert payload["issue"] in {"corrupt", "unreadable"}
    assert any(
        kw in payload["message"].lower()
        for kw in ["打开失败", "损坏", "failed", "无法读取", "cannot read", "no valid workbook"]
    )


def test_mcp_detect_corrupt_unsupported_legacy_xls(tmp_path: Path) -> None:
    legacy = tmp_path / "legacy.xls"
    legacy.write_bytes(b"BIFF\xff\xfe garbage that openpyxl cannot parse")

    payload = _call_tool("detect_corrupt_workbook", {"path": str(legacy)})

    assert payload["ok"] is False
    assert payload["issue"] == "unsupported"
    assert "不支持" in payload["message"] or "unsupported" in payload["message"].lower()


def test_mcp_detect_encoding_empty_file(tmp_path: Path) -> None:
    empty = tmp_path / "empty.csv"
    empty.write_bytes(b"")

    payload = _call_tool("detect_encoding", {"path": str(empty)})

    assert payload["encoding"] == "utf-8"
    assert payload["confidence"] == 0.0
    assert "空文件" in payload["message"] or "empty" in payload["message"].lower()


def test_mcp_detect_encoding_missing_file(tmp_path: Path) -> None:
    missing = tmp_path / "absent.csv"

    payload = _call_tool("detect_encoding", {"path": str(missing)})

    assert payload["encoding"] == "unknown"
    assert payload["confidence"] == 0.0
    assert "不存在" in payload["message"] or "not found" in payload["message"].lower()


def test_mcp_chunk_rejects_max_rows_below_one(tmp_path: Path) -> None:
    good = _make_xlsx(tmp_path / "ok.xlsx")

    payload = _call_tool(
        "chunk_large_workbook", {"path": str(good), "max_rows": 0}
    )

    assert payload["needs_chunking"] is False
    assert payload["chunks"] == []
    assert "max_rows" in payload["message"]


def test_mcp_chunk_corrupt_workbook_reports_open_failure(tmp_path: Path) -> None:
    bad = tmp_path / "broken.xlsx"
    bad.write_bytes(b"not a zip at all")

    payload = _call_tool("chunk_large_workbook", {"path": str(bad), "max_rows": 10})

    assert payload["needs_chunking"] is False
    assert payload["total_rows"] == 0
    assert payload["chunks"] == []
    assert "无法打开" in payload["message"] or "Cannot open" in payload["message"]


def test_mcp_chunk_missing_file(tmp_path: Path) -> None:
    missing = tmp_path / "absent.xlsx"

    payload = _call_tool(
        "chunk_large_workbook", {"path": str(missing), "max_rows": 10}
    )

    assert payload["needs_chunking"] is False
    assert payload["chunks"] == []
    assert "不存在" in payload["message"] or "not found" in payload["message"].lower()


def test_mcp_symlink_escaping_workspace_is_denied(tmp_path: Path) -> None:
    """A symlink inside the workspace pointing outside must be rejected.

    ``Path.resolve()`` follows the link before the sandbox check, so this
    surfaces as the generic path-denied message rather than the dedicated
    symlink branch.
    """
    fd, target = tempfile.mkstemp(prefix="excel-guard-symlink-target-", suffix=".xlsx")
    os.close(fd)
    try:
        link = tmp_path / "escape.xlsx"
        link.symlink_to(target)

        payload = _call_tool("detect_corrupt_workbook", {"path": str(link)})

        assert payload["ok"] is False
        assert payload["issue"] == "path_denied"
        assert "路径越界" in payload["message"] or "符号链接" in payload["message"]
    finally:
        Path(target).unlink(missing_ok=True)


def test_mcp_detect_encoding_high_confidence_non_utf8(tmp_path: Path) -> None:
    """UTF-16 sample yields high-confidence transcode hint (chardet is certain)."""
    csv_path = tmp_path / "utf16_big.csv"
    lines = ["name,city,amount", "Alice,Beijing,1000", "Bob,Shanghai,2000"]
    sample = "\n".join(lines * 50).encode("utf-16")
    csv_path.write_bytes(sample)

    payload = _call_tool("detect_encoding", {"path": str(csv_path)})

    assert payload["encoding"].lower().replace("_", "-") not in {"utf-8", "utf8"}
    assert payload["confidence"] >= 0.5
    assert "转码" in payload["message"] or "UTF-8" in payload["message"]


def test_mcp_call_unknown_tool_raises() -> None:
    with pytest.raises(Exception, match="Unknown tool"):
        _call_tool("does_not_exist", {"path": "x"})


def test_mcp_describe_workbook_returns_real_structure(tmp_path: Path) -> None:
    book = tmp_path / "shunde.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "顺德收支"
    ws.append(["交易日期", "金额", "收款行"])
    ws.append(["20260609", 4156.51, "顺德银行容桂支行"])
    wb.save(book)

    payload = _call_tool("describe_workbook", {"path": str(book)})

    assert payload["ok"] is True
    sheet = payload["sheets"][0]
    assert sheet["name"] == "顺德收支"
    assert sheet["row_count"] == 2
    assert sheet["column_count"] == 3
    assert sheet["headers"] == ["交易日期", "金额", "收款行"]
    preview = " ".join(" ".join(str(c) for c in row) for row in sheet["preview"])
    assert "顺德银行容桂支行" in preview


def test_mcp_sheet_to_markdown_range_excludes_outside_rows(tmp_path: Path) -> None:
    book = tmp_path / "range.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["交易日期", "金额", "备注"])
    ws.append(["20260101", 10, "INRANGE_TOKEN_AAA"])
    ws.append(["20260102", 20, "OUTRANGE_TOKEN_BBB"])
    wb.save(book)

    payload = _call_tool(
        "sheet_to_markdown",
        {"path": str(book), "sheet": "明细", "start_row": 2, "end_row": 2},
    )

    assert payload["ok"] is True
    assert "INRANGE_TOKEN_AAA" in payload["markdown"]
    assert "交易日期" in payload["markdown"]
    assert "OUTRANGE_TOKEN_BBB" not in payload["markdown"]


# ---------------------------------------------------------------------------
# 3. Server lifecycle
# ---------------------------------------------------------------------------


def test_tool_invocations_are_stateless_across_calls(tmp_path: Path) -> None:
    """Repeated calls must not share registry/invocation state."""
    good = _make_xlsx(tmp_path / "repeat.xlsx")

    first = _call_tool("detect_corrupt_workbook", {"path": str(good)})
    second = _call_tool("detect_corrupt_workbook", {"path": str(good)})

    assert first == second
    assert first["ok"] is True


def test_main_is_callable_and_invokes_mcp_run(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[tuple[tuple, dict]] = []

    def _fake_run(*args: Any, **kwargs: Any) -> None:
        calls.append((args, kwargs))

    monkeypatch.setattr(mcp, "run", _fake_run)

    main()

    assert calls == [((), {})]


def test_importing_server_module_does_not_start_stdio() -> None:
    """Import side effects must not block on a stdio transport.

    If module import ever called ``mcp.run()`` the subprocess would hang
    reading stdin and ``run`` would raise ``TimeoutExpired``.
    """
    env = {**os.environ, "PYTHONPATH": str(REPO_ROOT)}
    completed = subprocess.run(
        [
            sys.executable,
            "-c",
            "from poc.excel_guard_mcp.server import mcp, main; "
            "assert mcp.name == 'excel-guard'; "
            "assert callable(main); "
            "print('imported-ok')",
        ],
        cwd=str(REPO_ROOT),
        env=env,
        capture_output=True,
        text=True,
        timeout=20,
    )

    assert completed.returncode == 0, completed.stderr
    assert "imported-ok" in completed.stdout


def test_stdio_entrypoint_exits_on_stdin_eof() -> None:
    """``python -m poc.excel_guard_mcp`` must start via main() and shut down
    cleanly when stdin reaches EOF (exercises the __main__ entry block)."""
    env = {**os.environ, "PYTHONPATH": str(REPO_ROOT)}
    completed = subprocess.run(
        [sys.executable, "-m", "poc.excel_guard_mcp"],
        cwd=str(REPO_ROOT),
        env=env,
        stdin=subprocess.DEVNULL,
        capture_output=True,
        text=True,
        timeout=20,
    )

    assert completed.returncode == 0, completed.stderr
