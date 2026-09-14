# -*- coding: utf-8 -*-
"""Real tests for excel_guard readers (describe/slice/audit) + KB spreadsheet routing."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from poc.excel_guard_mcp.readers import (
    audit_workbook,
    describe_workbook,
    sheet_to_markdown,
)
from poc.kb_mcp.ingest import ingest_spreadsheet
from poc.kb_mcp.retrieve import search_knowledge
from poc.kb_mcp.spreadsheet import parse_spreadsheet_metadata
from poc.kb_mcp.stores import open_stores

REPO_ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(autouse=True)
def _sandbox(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """Confine path checks to the per-test temp dir; force local hash embeddings."""
    monkeypatch.setenv("POC_WORKSPACE", str(tmp_path))
    monkeypatch.setenv("POC_LOAD_SECRETS", "0")
    for key in (
        "POC_EMBEDDING_URL",
        "POC_EMBEDDING_API_KEY",
        "POC_CHAT_MODEL",
        "POC_CHAT_URL",
        "ARK_API_KEY",
        "ARK_BASE_URL",
        "ELASTICSEARCH_URL",
        "MYSQL_URL",
        "GALASYBASE_URL",
        "POC_ES_URL",
        "POC_MYSQL_URL",
        "POC_GALASYBASE_URL",
        "MINERU_ENDPOINT",
    ):
        monkeypatch.delenv(key, raising=False)
    return tmp_path


def _write_formula_workbook(path: Path) -> Path:
    """Title row + header row + a repeated formula column (uncached values)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "收入表"
    ws["A1"] = "2024年度汇总"
    ws.merge_cells("A1:D1")
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "区域", "数量", "单价", "收入"
    rows = [("北方", 1200, 9.99), ("南方", 840, 9.99), ("西部", 1500, 9.99)]
    for i, (region, qty, price) in enumerate(rows, start=3):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = region, qty, price
        ws[f"D{i}"] = f"=B{i}*C{i}"
    wb.save(path)
    return path


def test_describe_workbook_votes_header_and_aliases_formulas(tmp_path: Path) -> None:
    xlsx = _write_formula_workbook(tmp_path / "收入.xlsx")

    result = describe_workbook(str(xlsx), preview_rows=4)

    assert result["ok"] is True
    assert result["kind"] == "xlsx"
    sheets = result["sheets"]
    assert [s["name"] for s in sheets] == ["收入表"]
    sheet = sheets[0]
    # header voting must skip the merged title row and pick row 2
    assert sheet["header_row"] == 2
    assert sheet["headers"][:4] == ["区域", "数量", "单价", "收入"]
    assert sheet["row_count"] == 5
    assert sheet["column_count"] == 4
    assert "A1:D1" in sheet["merged_ranges"]
    # the repeated =B*C column collapses into one alias pattern
    aliases = sheet["formula_aliases"]
    assert len(aliases) == 1
    assert aliases[0]["pattern"] == "=B1*C1"
    assert aliases[0]["count"] == 3
    # preview substitutes formulas for uncached cells instead of blanks
    assert "=B3*C3" in sheet["markdown"]
    assert sheet["note"]


def test_describe_workbook_csv_reports_encoding_and_rows(tmp_path: Path) -> None:
    csv_path = tmp_path / "订单.csv"
    csv_path.write_bytes("编号,城市\n1,北京\n2,广州\n".encode("gbk"))

    result = describe_workbook(str(csv_path), preview_rows=2)

    assert result["ok"] is True
    assert result["kind"] == "csv"
    sheet = result["sheets"][0]
    assert sheet["headers"] == ["编号", "城市"]
    assert sheet["row_count"] == 3
    assert "gb" in sheet["note"].lower() or "utf" in sheet["note"].lower()
    assert "| 编号 | 城市 |" in sheet["markdown"]


def test_sheet_to_markdown_returns_exact_row_range(tmp_path: Path) -> None:
    xlsx = _write_formula_workbook(tmp_path / "收入.xlsx")

    result = sheet_to_markdown(str(xlsx), start_row=4, end_row=5)

    assert result["ok"] is True
    assert result["n_rows"] == 2
    assert "| 南方 | 840 |" in result["markdown"]
    assert "| 西部 | 1500 |" in result["markdown"]
    assert "北方" not in result["markdown"]

    missing = sheet_to_markdown(str(xlsx), sheet="不存在", start_row=1, end_row=2)
    assert missing["ok"] is False
    assert "收入表" in missing["message"]


def test_sheet_to_markdown_rejects_oversize_range(tmp_path: Path) -> None:
    xlsx = _write_formula_workbook(tmp_path / "收入.xlsx")

    result = sheet_to_markdown(str(xlsx), start_row=1, end_row=5000)

    assert result["ok"] is False
    assert "分块" in result["message"] or "chunk" in result["message"].lower()


def _write_error_workbook(path: Path) -> Path:
    """Handcraft an xlsx with cached error / literal-in-formula cells.

    openpyxl cannot compute formulas, so the cached #DIV/0! is written at the
    XML level to exercise the audit's value pass for real (no mocks).
    """
    sheet_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        "<sheetData>"
        '<row r="1">'
        '<c r="A1"><v>10</v></c>'
        '<c r="B1" t="e"><f>A1/0</f><v>#DIV/0!</v></c>'
        '<c r="C1"><f>SUM(1000,A1)</f><v>1010</v></c>'
        "</row>"
        "</sheetData>"
        "</worksheet>"
    )
    workbook_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets>'
        "</workbook>"
    )
    workbook_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
        'Target="worksheets/sheet1.xml"/>'
        "</Relationships>"
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        "</Types>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", root_rels)
        zf.writestr("xl/workbook.xml", workbook_xml)
        zf.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        zf.writestr("xl/worksheets/sheet1.xml", sheet_xml)
    return path


def test_audit_workbook_classifies_errors_and_literals(tmp_path: Path) -> None:
    xlsx = _write_error_workbook(tmp_path / "错误表.xlsx")

    result = audit_workbook(str(xlsx))

    assert result["ok"] is True
    assert result["n_formula_cells"] == 2
    must_fix = result["must_fix"]
    assert len(must_fix) == 1
    assert must_fix[0]["cell"] == "B1"
    assert must_fix[0]["error"] == "#DIV/0!"
    assert must_fix[0]["formula"] == "=A1/0"
    review = result["review"]
    assert len(review) == 1
    assert review[0]["cell"] == "C1"
    assert "1000" in review[0]["formula"]


def test_audit_workbook_skips_csv(tmp_path: Path) -> None:
    csv_path = tmp_path / "plain.csv"
    csv_path.write_text("a,b\n1,2\n", encoding="utf-8")

    result = audit_workbook(str(csv_path))

    assert result["ok"] is True
    assert result["n_formula_cells"] == 0
    assert result["must_fix"] == []


def test_spreadsheet_metadata_routes_through_kb(tmp_path: Path) -> None:
    wb = Workbook()
    sales = wb.active
    sales.title = "销售明细"
    sales.append(["日期", "地区", "销售额"])
    sales.append(["2024-01-05", "顺德", 52000])
    people = wb.create_sheet("人员名单")
    people.append(["姓名", "部门"])
    people.append(["张三", "零售部"])
    xlsx = tmp_path / "经营数据.xlsx"
    wb.save(xlsx)

    parsed = parse_spreadsheet_metadata(str(xlsx))

    assert parsed["ok"] is True
    assert parsed["doc_id"] == "经营数据"
    assert parsed["sheets"] == ["销售明细", "人员名单"]
    assert len(parsed["chunks"]) == 2
    sales_chunk = parsed["chunks"][0]
    assert sales_chunk["kind"] == "sheet_meta"
    assert "销售明细" in sales_chunk["text"]
    assert "销售额" in sales_chunk["text"]
    assert sales_chunk["next_id"].endswith("p2:sheet_meta:0")

    bundle = open_stores(root=tmp_path / "kb_store")
    result = ingest_spreadsheet(str(xlsx), stores=bundle)

    assert result["ok"] is True
    assert result["n_chunks"] == 2
    assert result["backend"] == "spreadsheet-meta"

    search = search_knowledge("销售明细的销售额", stores=bundle)

    assert search["ok"] is True
    assert search["hits"]
    top_kinds = {hit.get("kind") for hit in search["hits"]}
    assert "sheet_meta" in top_kinds
    top = search["hits"][0]
    assert top["doc_id"] == "经营数据"
    assert top["chapter"] == "销售明细"
