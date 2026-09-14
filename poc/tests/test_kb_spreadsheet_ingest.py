# -*- coding: utf-8 -*-
"""Spreadsheet metadata ingest: route by headers/sheets/samples, never dump data rows."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from poc.excel_guard_mcp.readers import describe_workbook
from poc.kb_mcp.ingest import ingest_spreadsheet
from poc.kb_mcp.retrieve import search_knowledge
from poc.kb_mcp.stores import open_stores

REPO_ROOT = Path(__file__).resolve().parents[2]

DEEP_TOKEN = "DEEP_CELL_TOKEN_NEVER_INDEX_7f3a"
SHUNDE_TOKEN = "ShundeMetaTokenX9"
JIAOZHOU_TOKEN = "JiaozhouMetaTokenY2"


@pytest.fixture(autouse=True)
def _sandbox(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setenv("POC_WORKSPACE", str(tmp_path))
    monkeypatch.setenv("POC_LOAD_SECRETS", "0")
    for key in (
        "MINERU_ENDPOINT",
        "POC_EMBEDDING_URL",
        "POC_EMBEDDING_API_KEY",
        "POC_CHAT_MODEL",
        "POC_CHAT_URL",
        "ARK_API_KEY",
        "ARK_BASE_URL",
        "ELASTICSEARCH_URL",
        "MYSQL_URL",
        "GALASYBASE_URL",
        "POC_ES_URL",
        "POC_MYSQL_URL",
        "POC_GALASYBASE_URL",
    ):
        monkeypatch.delenv(key, raising=False)
    return tmp_path


def _write_shunde(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "顺德收支台账"
    ws.append(["交易日期", "金额", f"顺德网点{SHUNDE_TOKEN}"])
    ws.append(["20260601", 100, "顺德容桂支行"])
    ws.append(["20260602", 110, "顺德大良支行"])
    ws.append(["20260603", 120, "顺德伦教支行"])
    ws.append(["20260604", 130, "顺德北滘支行"])
    ws.append(["20260605", 140, "顺德陈村支行"])
    for i in range(14):
        ws.append([f"202607{i+1:02d}", i, "padding-row"])
    ws.append(["20260820", 999, DEEP_TOKEN])
    wb.save(path)
    return path


def _write_jiaozhou(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "胶州收支台账"
    ws.append(["交易日期", "金额", f"胶州网点{JIAOZHOU_TOKEN}"])
    ws.append(["20260601", 50, "胶州铺集支行"])
    wb.save(path)
    return path


def test_ingest_spreadsheet_metadata_routes_to_matching_file(
    tmp_path: Path,
) -> None:
    shunde = _write_shunde(tmp_path / "shunde_ledger.xlsx")
    other = _write_jiaozhou(tmp_path / "jiaozhou_ledger.xlsx")

    first = ingest_spreadsheet(str(shunde), doc_id="shunde-ledger", sample_rows=5)
    second = ingest_spreadsheet(str(other), doc_id="jiaozhou-ledger", sample_rows=5)
    assert first["ok"] is True
    assert second["ok"] is True
    assert first["n_chunks"] >= 1

    hits = search_knowledge(SHUNDE_TOKEN, k=8)
    assert hits["ok"] is True
    assert hits["hits"], hits
    top = hits["hits"][0]
    assert top["doc_id"] == "shunde-ledger"
    assert SHUNDE_TOKEN in (top.get("text") or "")
    assert "顺德收支台账" in (top.get("text") or top.get("chapter") or "")
    assert top["doc_id"] != "jiaozhou-ledger"
    assert JIAOZHOU_TOKEN not in (top.get("text") or "")


def test_ingest_spreadsheet_does_not_index_deep_data_cells(tmp_path: Path) -> None:
    shunde = _write_shunde(tmp_path / "shunde_ledger.xlsx")
    result = ingest_spreadsheet(str(shunde), doc_id="shunde-ledger", sample_rows=5)
    assert result["ok"] is True

    indexed = " ".join(
        (row.get("text") or "")
        for row in open_stores().relational.list_chunks(doc_id="shunde-ledger")
    )
    assert SHUNDE_TOKEN in indexed
    assert "顺德容桂支行" in indexed
    assert DEEP_TOKEN not in indexed
    assert "padding-row" not in indexed


def test_ingest_spreadsheet_denies_outside_workspace(
    tmp_path: Path, tmp_path_factory: pytest.TempPathFactory
) -> None:
    secret_dir = tmp_path_factory.mktemp("outside-kb-sheet")
    secret = secret_dir / "secret_sheet_outside.xlsx"
    _write_shunde(secret)

    result = ingest_spreadsheet(str(secret), doc_id="leaky")
    assert result["ok"] is False
    assert result.get("issue") == "path_denied" or "越界" in (result.get("message") or "")
    assert "secret_sheet_outside" not in (result.get("message") or "")
    assert DEEP_TOKEN not in (result.get("message") or "")
    assert open_stores().relational.list_chunks(doc_id="leaky") == []


def test_shunde_income_fixture_is_metadata_only(tmp_path: Path) -> None:
    src = REPO_ROOT / "poc" / "fixtures" / "收支明细_顺德.xlsx"
    assert src.is_file()
    dest = tmp_path / src.name
    dest.write_bytes(src.read_bytes())

    described = describe_workbook(str(dest), preview_rows=5)
    assert described["ok"] is True
    ingested = ingest_spreadsheet(str(dest), doc_id="shunde-income", sample_rows=5)
    assert ingested["ok"] is True
    assert ingested["n_chunks"] == 1

    from openpyxl import load_workbook

    wb = load_workbook(dest, read_only=True, data_only=True)
    ws = wb.active
    header = None
    deep_value = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            header = ["" if c is None else str(c) for c in row]
            continue
        if i == 200:
            cells = ["" if c is None else str(c) for c in row]
            # PTA_EFTNO is unique per row; fall back to last non-empty cell.
            if header and "PTA_EFTNO" in header:
                deep_value = cells[header.index("PTA_EFTNO")]
            else:
                deep_value = next((c for c in reversed(cells) if c), None)
            break
    wb.close()
    assert deep_value

    indexed = " ".join(
        row.get("text") or ""
        for row in open_stores().relational.list_chunks(doc_id="shunde-income")
    )
    assert "收款行" in indexed
    assert "顺德" in indexed
    assert "青岛" not in indexed
    assert deep_value not in indexed

    hits = search_knowledge("顺德银行", k=4)
    assert hits["ok"] is True
    assert hits["hits"]
    assert hits["hits"][0]["doc_id"] == "shunde-income"
