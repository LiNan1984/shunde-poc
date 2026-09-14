# -*- coding: utf-8 -*-
"""Real tests for excel_guard_mcp.guards (no mocks of logic under test)."""

from __future__ import annotations

import os
import stat
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from poc.excel_guard_mcp import guards as guards_mod
from poc.excel_guard_mcp.guards import (
    _ENCODING_HARD_LIMIT,
    chunk_large_workbook,
    detect_corrupt_workbook,
    detect_encoding,
)
from poc.excel_guard_mcp.readers import describe_workbook, sheet_to_markdown

REPO_ROOT = Path(__file__).resolve().parents[2]
FIXTURES = REPO_ROOT / "poc" / "fixtures"


@pytest.fixture(autouse=True)
def _sandbox_workspace(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Confine path checks to the per-test temp dir (plus allow fixture tests)."""
    monkeypatch.setenv("POC_WORKSPACE", str(tmp_path))


def test_detect_corrupt_workbook_rejects_invalid_zip(tmp_path: Path) -> None:
    bad = tmp_path / "corrupt.xlsx"
    bad.write_bytes(b"this is not a valid zip or xlsx file")

    result = detect_corrupt_workbook(str(bad))

    assert result["ok"] is False
    assert result["issue"] == "corrupt"
    assert "损坏" in result["message"] or "Corrupt" in result["message"]


def test_detect_corrupt_workbook_accepts_valid_xlsx(tmp_path: Path) -> None:
    good = tmp_path / "ok.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "hello"
    ws["B1"] = 42
    wb.save(good)

    result = detect_corrupt_workbook(str(good))

    assert result["ok"] is True
    assert result["issue"] == ""
    assert "OK" in result["message"] or "正常" in result["message"]


def test_detect_encoding_finds_latin1_csv(tmp_path: Path) -> None:
    csv_path = tmp_path / "latin1.csv"
    csv_path.write_bytes("name,city\nAlice,Caf\xe9\n".encode("latin-1"))

    result = detect_encoding(str(csv_path))

    encoding = result["encoding"].lower().replace("_", "-")
    assert encoding not in {"utf-8", "utf8"}
    assert "UTF-8" in result["message"] or "utf-8" in result["message"].lower()
    assert "编码" in result["message"] or "encoding" in result["message"].lower()


def test_detect_encoding_xlsx_notes_zip_utf8(tmp_path: Path) -> None:
    xlsx = tmp_path / "note.xlsx"
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(xlsx)

    result = detect_encoding(str(xlsx))

    assert result["encoding"].lower() in {"utf-8", "utf8"}
    assert "ZIP" in result["message"] or "zip" in result["message"]


def test_chunk_large_workbook_needs_chunking(tmp_path: Path) -> None:
    large = tmp_path / "large.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["id", "value"])
    for i in range(120):
        ws.append([i, f"v{i}"])
    wb.save(large)

    result = chunk_large_workbook(str(large), max_rows=50)

    assert result["needs_chunking"] is True
    assert result["total_rows"] > 50
    assert len(result["chunks"]) >= 2
    first = result["chunks"][0]
    assert first["start_row"] == 1
    assert first["end_row"] == 50
    assert first["sheet"] == "Sheet1"
    assert "分块" in result["message"] or "chunk" in result["message"].lower()


def test_chunk_large_workbook_no_chunking_when_small(tmp_path: Path) -> None:
    small = tmp_path / "small.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["a", "b"])
    ws.append([1, 2])
    wb.save(small)

    result = chunk_large_workbook(str(small), max_rows=5000)

    assert result["needs_chunking"] is False
    assert result["total_rows"] <= 5000
    assert result["chunks"] == []
    assert "无需分块" in result["message"] or "No chunking" in result["message"]


def test_path_outside_workspace_denied(tmp_path: Path) -> None:
    outside = Path("/etc/passwd")
    if not outside.exists():
        pytest.skip("/etc/passwd not present")

    result = detect_corrupt_workbook(str(outside))
    assert result["ok"] is False
    assert result["issue"] == "path_denied"


def test_committed_fixtures_smoke(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("POC_WORKSPACE", str(REPO_ROOT))

    corrupt = detect_corrupt_workbook(str(FIXTURES / "corrupt.xlsx"))
    assert corrupt["ok"] is False
    assert corrupt["issue"] == "corrupt"

    enc_path = FIXTURES / "encoding_gbk.csv"
    if not enc_path.exists():
        enc_path = FIXTURES / "encoding_latin1.csv"
    enc = detect_encoding(str(enc_path))
    assert enc["encoding"].lower() not in {"utf-8", "utf8"}

    large = chunk_large_workbook(str(FIXTURES / "large_chunk_demo.xlsx"), max_rows=5000)
    assert large["needs_chunking"] is True
    assert large["total_rows"] > 5000

    income = FIXTURES / "收支明细_顺德.xlsx"
    assert income.is_file()
    assert "测试集" not in income.name
    described = describe_workbook(str(income))
    assert described["ok"] is True
    sheet = described["sheets"][0]
    assert sheet["name"] == "Sheet1"
    assert sheet["row_count"] == 267
    assert sheet["column_count"] == 14
    assert "收款行" in sheet["headers"]
    assert "交易日期" in sheet["headers"]
    preview_blob = " ".join(" ".join(str(c) for c in row) for row in sheet["preview"])
    assert "顺德" in preview_blob
    assert "青岛" not in preview_blob


# ---------------------------------------------------------------------------
# Sandbox / security edge cases
# ---------------------------------------------------------------------------


def test_symlink_escape_outside_workspace_denied(
    tmp_path: Path, tmp_path_factory: pytest.TempPathFactory
) -> None:
    secret_dir = tmp_path_factory.mktemp("outside")
    secret = secret_dir / "secret.csv"
    secret.write_text("id,name\n1,alice\n", encoding="utf-8")

    link = tmp_path / "link.csv"
    link.symlink_to(secret)

    result = detect_encoding(str(link))

    assert result["encoding"] == "unknown"
    assert result["confidence"] == 0.0
    assert "越界" in result["message"] or "outside" in result["message"]


def test_broken_symlink_target_outside_workspace_denied(
    tmp_path: Path, tmp_path_factory: pytest.TempPathFactory
) -> None:
    outside_dir = tmp_path_factory.mktemp("outside2")
    link = tmp_path / "dangling.csv"
    link.symlink_to(outside_dir / "does_not_exist.csv")

    result = detect_encoding(str(link))

    assert result["encoding"] == "unknown"
    assert result["confidence"] == 0.0
    assert "越界" in result["message"] or "outside" in result["message"]


def test_symlink_loop_reported_as_invalid(tmp_path: Path) -> None:
    loop = tmp_path / "loop.xlsx"
    loop.symlink_to(loop.name)  # self-referencing symlink (ELOOP)

    result = detect_corrupt_workbook(str(loop))

    assert result["ok"] is False
    assert result["issue"] == "path_denied"
    assert "符号链接" in result["message"]


def test_relative_path_traversal_is_denied(tmp_path: Path) -> None:
    # Deterministic traversal anchored inside the workspace.
    escape = tmp_path / "sub" / ".." / ".." / "evil.csv"

    result = detect_encoding(str(escape))

    assert result["encoding"] == "unknown"
    assert result["confidence"] == 0.0
    assert "越界" in result["message"] or "outside" in result["message"]

    # Literal cwd-relative traversal must also be denied.
    literal = detect_corrupt_workbook("../../../etc/passwd")
    assert literal["ok"] is False
    assert literal["issue"] == "path_denied"


def test_path_with_dotdot_inside_workspace_is_allowed(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    (data_dir / "sub").mkdir(parents=True)
    note = data_dir / "note.csv"
    note.write_text("a,b\n1,2\n", encoding="ascii")

    traversed = tmp_path / "data" / "sub" / ".." / "note.csv"
    result = detect_encoding(str(traversed))

    assert result["encoding"] == "ascii"
    assert "越界" not in result["message"]


def test_missing_workspace_env_falls_back_to_repo_root(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("POC_WORKSPACE", raising=False)
    monkeypatch.delenv("QWENPAW_WORKING_DIR", raising=False)

    # The committed fixture lives under the repository root, so it is allowed.
    result = detect_corrupt_workbook(str(FIXTURES / "corrupt.xlsx"))

    assert result["ok"] is False
    assert result["issue"] == "corrupt"

    # A path outside the repository is still denied.
    outside = Path("/etc/passwd")
    if outside.exists():
        denied = detect_encoding(str(outside))
        assert denied["encoding"] == "unknown"
        assert "越界" in denied["message"] or "outside" in denied["message"]


def test_qwenpaw_working_dir_env_fallback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv("POC_WORKSPACE", raising=False)
    monkeypatch.setenv("QWENPAW_WORKING_DIR", str(tmp_path))

    good = tmp_path / "ok.xlsx"
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(good)

    result = detect_corrupt_workbook(str(good))

    assert result["ok"] is True


def test_unresolvable_path_permission_error_denied(tmp_path: Path) -> None:
    if os.geteuid() == 0:
        pytest.skip("root bypasses directory permissions")

    locked = tmp_path / "locked"
    locked.mkdir()
    target = locked / "f.xlsx"
    os.chmod(locked, 0o000)
    try:
        result = detect_corrupt_workbook(str(target))
    except PermissionError:
        pytest.skip("Path.resolve does not trap permission errors on this platform")
    finally:
        os.chmod(locked, stat.S_IRWXU)

    assert result["ok"] is False
    assert result["issue"] == "path_denied"
    assert "路径" in result["message"] and ("denied" in result["message"].lower() or "无法" in result["message"])


def test_nonexistent_file_inside_workspace(tmp_path: Path) -> None:
    missing = tmp_path / "missing.xlsx"

    corrupt = detect_corrupt_workbook(str(missing))
    assert corrupt["ok"] is False
    assert corrupt["issue"] == "missing"
    assert "不存在" in corrupt["message"]

    encoding = detect_encoding(str(missing))
    assert encoding["encoding"] == "unknown"
    assert encoding["confidence"] == 0.0
    assert "不存在" in encoding["message"]

    chunk = chunk_large_workbook(str(missing))
    assert chunk["needs_chunking"] is False
    assert chunk["total_rows"] == 0
    assert chunk["chunks"] == []
    assert "不存在" in chunk["message"]


# ---------------------------------------------------------------------------
# detect_corrupt_workbook edge cases
# ---------------------------------------------------------------------------


def test_detect_corrupt_zip_without_content_types(tmp_path: Path) -> None:
    fake = tmp_path / "no_content_types.xlsx"
    with zipfile.ZipFile(fake, "w") as zf:
        zf.writestr("hello.txt", "not an office package")

    result = detect_corrupt_workbook(str(fake))

    assert result["ok"] is False
    assert result["issue"] == "corrupt"
    assert "[Content_Types].xml" in result["message"]


def test_detect_corrupt_old_xls_format_unsupported(tmp_path: Path) -> None:
    old = tmp_path / "legacy.xls"
    # OLE2 compound document magic header followed by junk (not a real .xls).
    old.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)

    result = detect_corrupt_workbook(str(old))

    assert result["ok"] is False
    assert result["issue"] == "unsupported"
    assert "不支持" in result["message"] or "unsupported" in result["message"].lower()


def test_detect_corrupt_empty_xlsx(tmp_path: Path) -> None:
    empty = tmp_path / "empty.xlsx"
    empty.write_bytes(b"")

    result = detect_corrupt_workbook(str(empty))

    assert result["ok"] is False
    assert result["issue"] == "corrupt"
    assert "ZIP" in result["message"]


def test_detect_corrupt_truncated_zip_header(tmp_path: Path) -> None:
    truncated = tmp_path / "truncated.xlsx"
    truncated.write_bytes(b"PK\x03\x04" + b"\x00" * 30)

    result = detect_corrupt_workbook(str(truncated))

    assert result["ok"] is False
    assert result["issue"] == "corrupt"
    assert "ZIP" in result["message"]


def test_detect_corrupt_zip_missing_workbook_parts(tmp_path: Path) -> None:
    # Valid ZIP with [Content_Types].xml but no xl/workbook.xml: openpyxl fails.
    fake = tmp_path / "skeleton.xlsx"
    with zipfile.ZipFile(fake, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("_rels/.rels", "<Relationships/>")

    result = detect_corrupt_workbook(str(fake))

    assert result["ok"] is False
    # Missing workbook parts: could be caught as corrupt (openpyxl) or unreadable (OSError)
    assert result["issue"] in {"corrupt", "unreadable"}
    assert any(
        kw in result["message"]
        for kw in ["打开失败", "损坏", "failed", "无法读取", "cannot read", "no valid workbook"]
    )


# ---------------------------------------------------------------------------
# detect_encoding edge cases
# ---------------------------------------------------------------------------


def test_detect_encoding_empty_file(tmp_path: Path) -> None:
    empty = tmp_path / "empty.csv"
    empty.write_bytes(b"")

    result = detect_encoding(str(empty))

    assert result["encoding"] == "utf-8"
    assert result["confidence"] == 0.0
    assert "空文件" in result["message"] or "Empty" in result["message"]


def test_detect_encoding_ascii_csv(tmp_path: Path) -> None:
    csv_path = tmp_path / "ascii.csv"
    csv_path.write_text("id,name\n1,Alice\n2,Bob\n", encoding="ascii")

    result = detect_encoding(str(csv_path))

    assert result["encoding"] == "ascii"
    assert result["confidence"] == 1.0
    assert "OK" in result["message"] or "正常" in result["message"]


def test_detect_encoding_gbk_chinese_csv(tmp_path: Path) -> None:
    csv_path = tmp_path / "gbk.csv"
    content = "编号,城市,备注说明\n1,北京,欢迎参加测试\n2,上海,研发与创新中心\n" * 4
    csv_path.write_bytes(content.encode("gbk"))

    result = detect_encoding(str(csv_path))

    assert result["encoding"].lower().startswith("gb")
    assert result["encoding"].lower() not in {"utf-8", "utf8"}
    assert "UTF-8" in result["message"]
    assert "转码" in result["message"]


def test_detect_encoding_file_over_hard_limit(tmp_path: Path) -> None:
    huge = tmp_path / "huge.csv"
    # Sparse file: seek past the 32 MiB limit without actually writing that much.
    with huge.open("wb") as fh:
        fh.seek(_ENCODING_HARD_LIMIT + 10)
        fh.write(b"\n")
    assert huge.stat().st_size > _ENCODING_HARD_LIMIT

    result = detect_encoding(str(huge))

    # After M4 fix: we sample the first 1 MiB instead of refusing entirely.
    # The result should not be "unknown" with 0.0 confidence (we did sample).
    # It should note that only a sample was used.
    assert "sampled" in result["message"].lower() or "采样" in result["message"]
    assert result["encoding"] != "unknown" or result["confidence"] > 0.0


def test_detect_encoding_low_confidence_short_file(tmp_path: Path) -> None:
    weird = tmp_path / "weird.csv"
    weird.write_bytes(b"\x80\x81\x82")

    result = detect_encoding(str(weird))

    assert result["encoding"].lower() not in {"utf-8", "utf8", "ascii"}
    assert result["confidence"] < 0.5
    assert "置信度偏低" in result["message"] or "Low-confidence" in result["message"]


def test_detect_encoding_utf16_high_confidence_non_utf8(tmp_path: Path) -> None:
    csv_path = tmp_path / "utf16.csv"
    payload = "名称,数值\n甲,100\n乙,200\n".encode("utf-16-le")
    csv_path.write_bytes(b"\xff\xfe" + payload)  # UTF-16 LE BOM

    result = detect_encoding(str(csv_path))

    assert result["encoding"].lower().replace("-", "") == "utf16"
    assert result["confidence"] >= 0.9
    assert "非 UTF-8" in result["message"]


def test_detect_encoding_invalid_bytes_labeled_utf8_falls_back(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # chardet is an external dependency; force its (rare) mislabel so the
    # defensive UnicodeDecodeError fallback path is exercised.
    bad = tmp_path / "mislabeled.csv"
    bad.write_bytes(b"col1,col2\n\xff\xfe\xfd\x80")
    monkeypatch.setattr(
        guards_mod.chardet,
        "detect",
        lambda raw: {"encoding": "utf-8", "confidence": 0.9, "language": ""},
    )

    result = detect_encoding(str(bad))

    assert result["encoding"] == "utf-8"
    assert result["confidence"] == 0.9
    assert "非 UTF-8" in result["message"]


def test_detect_encoding_corrupt_xlsx_reports_not_zip(tmp_path: Path) -> None:
    bad = tmp_path / "broken.xlsx"
    bad.write_bytes(b"definitely not a zip")

    result = detect_encoding(str(bad))

    assert result["encoding"] == "utf-8"
    assert result["confidence"] == 0.0
    assert "不是有效 ZIP" in result["message"]


# ---------------------------------------------------------------------------
# chunk_large_workbook edge cases
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("bad_max_rows", [0, -1, -100])
def test_chunk_rejects_nonpositive_max_rows(
    tmp_path: Path, bad_max_rows: int
) -> None:
    small = tmp_path / "small.xlsx"
    wb = Workbook()
    wb.active.append(["a"])
    wb.save(small)

    result = chunk_large_workbook(str(small), max_rows=bad_max_rows)

    assert result["needs_chunking"] is False
    assert result["total_rows"] == 0
    assert result["chunks"] == []
    assert "max_rows" in result["message"]


def test_chunk_max_rows_one_boundary(tmp_path: Path) -> None:
    book = tmp_path / "three_rows.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Solo"
    for i in range(3):
        ws.append([i])
    wb.save(book)

    result = chunk_large_workbook(str(book), max_rows=1)

    assert result["needs_chunking"] is True
    assert result["total_rows"] == 3
    assert result["chunks"] == [
        {"start_row": 1, "end_row": 1, "sheet": "Solo"},
        {"start_row": 2, "end_row": 2, "sheet": "Solo"},
        {"start_row": 3, "end_row": 3, "sheet": "Solo"},
    ]


def test_chunk_multiple_sheets_mixed(tmp_path: Path) -> None:
    book = tmp_path / "mixed.xlsx"
    wb = Workbook()
    small = wb.active
    small.title = "Small"
    small.append(["k"])
    small.append(["v"])
    big = wb.create_sheet("Big")
    big.append(["id"])
    for i in range(120):
        big.append([i])
    wb.save(book)

    result = chunk_large_workbook(str(book), max_rows=50)

    assert result["needs_chunking"] is True
    assert result["total_rows"] == 121
    assert len(result["chunks"]) == 3
    assert all(c["sheet"] == "Big" for c in result["chunks"])
    assert result["chunks"][0]["start_row"] == 1
    assert result["chunks"][0]["end_row"] == 50
    assert result["chunks"][-1]["end_row"] == 121


def test_chunk_exact_boundary_no_chunking(tmp_path: Path) -> None:
    book = tmp_path / "exact.xlsx"
    wb = Workbook()
    ws = wb.active
    for i in range(10):
        ws.append([i])
    wb.save(book)

    result = chunk_large_workbook(str(book), max_rows=10)

    assert result["needs_chunking"] is False
    assert result["total_rows"] == 10
    assert result["chunks"] == []
    assert "无需分块" in result["message"] or "No chunking" in result["message"]


def test_chunk_corrupt_workbook_graceful_error(tmp_path: Path) -> None:
    bad = tmp_path / "corrupt.xlsx"
    bad.write_bytes(b"not an excel file at all")

    result = chunk_large_workbook(str(bad), max_rows=100)

    assert result["needs_chunking"] is False
    assert result["total_rows"] == 0
    assert result["chunks"] == []
    assert "无法打开工作簿" in result["message"] or "Cannot open" in result["message"]


# ---------------------------------------------------------------------------
# Integration: consistent sandbox enforcement + extension support
# ---------------------------------------------------------------------------


def test_all_guards_consistently_enforce_workspace_sandbox() -> None:
    outside = Path("/etc/passwd")
    if not outside.exists():
        pytest.skip("/etc/passwd not present")

    corrupt = detect_corrupt_workbook(str(outside))
    assert corrupt["ok"] is False
    assert corrupt["issue"] == "path_denied"

    encoding = detect_encoding(str(outside))
    assert encoding["encoding"] == "unknown"
    assert encoding["confidence"] == 0.0
    assert "越界" in encoding["message"] or "outside" in encoding["message"]

    chunk = chunk_large_workbook(str(outside))
    assert chunk["needs_chunking"] is False
    assert chunk["total_rows"] == 0
    assert chunk["chunks"] == []
    assert "越界" in chunk["message"] or "outside" in chunk["message"]

    described = describe_workbook(str(outside))
    assert described["ok"] is False
    assert described["issue"] == "path_denied"
    assert "越界" in described["message"] or "outside" in described["message"]
    assert outside.name not in described["message"]

    markdown = sheet_to_markdown(str(outside), start_row=1, end_row=2)
    assert markdown["ok"] is False
    assert markdown["issue"] == "path_denied"
    assert "越界" in markdown["message"] or "outside" in markdown["message"]
    assert outside.name not in markdown["message"]


@pytest.mark.parametrize("extension", [".xlsm", ".xltx"])
def test_macro_and_template_extensions_accepted(
    tmp_path: Path, extension: str
) -> None:
    book = tmp_path / f"book{extension}"
    wb = Workbook()
    wb.active["A1"] = "ok"
    wb.save(book)

    corrupt = detect_corrupt_workbook(str(book))
    assert corrupt["ok"] is True

    encoding = detect_encoding(str(book))
    assert encoding["encoding"].lower() in {"utf-8", "utf8"}
    assert encoding["confidence"] == 1.0
    assert "ZIP" in encoding["message"]


def test_server_module_imports() -> None:
    from poc.excel_guard_mcp.server import mcp

    assert mcp is not None
    assert getattr(mcp, "name", None) or True


# ---------------------------------------------------------------------------
# describe_workbook / sheet_to_markdown
# ---------------------------------------------------------------------------


def _write_ledger(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "顺德收支"
    ws.append(["交易日期", "金额", "收款行"])
    ws.append(["20260609", 4156.51, "顺德银行容桂支行"])
    ws.append(["20260610", 200.0, "OUTRANGE_TOKEN_BBB"])
    ws.append(["20260611", 300.0, "DEEP_TOKEN_CCC"])
    extra = wb.create_sheet("网点清单")
    extra.append(["网点", "城市"])
    extra.append(["容桂", "顺德"])
    wb.save(path)
    return path


def test_describe_workbook_returns_real_sheet_structure(tmp_path: Path) -> None:
    book = _write_ledger(tmp_path / "shunde_ledger.xlsx")

    result = describe_workbook(str(book), preview_rows=1)

    assert result["ok"] is True
    assert result.get("issue") in {"", None}
    sheets = {s["name"]: s for s in result["sheets"]}
    assert set(sheets) == {"顺德收支", "网点清单"}
    detail = sheets["顺德收支"]
    assert detail["row_count"] == 4
    assert detail["column_count"] == 3
    assert detail["headers"] == ["交易日期", "金额", "收款行"]
    preview_blob = " ".join(
        " ".join(str(c) for c in row) for row in detail["preview"]
    )
    assert "顺德银行容桂支行" in preview_blob
    assert "交易日期" in preview_blob
    # preview_rows=1 → header + first data row only
    assert "OUTRANGE_TOKEN_BBB" not in preview_blob
    assert "DEEP_TOKEN_CCC" not in preview_blob
    branch = sheets["网点清单"]
    assert branch["headers"] == ["网点", "城市"]
    assert branch["row_count"] == 2


def test_sheet_to_markdown_keeps_header_and_only_in_range_rows(
    tmp_path: Path,
) -> None:
    book = _write_ledger(tmp_path / "range.xlsx")

    result = sheet_to_markdown(
        str(book), sheet="顺德收支", start_row=2, end_row=2
    )

    assert result["ok"] is True
    md = result["markdown"]
    assert result["sheet"] == "顺德收支"
    assert "交易日期" in md
    assert "收款行" in md
    assert "顺德银行容桂支行" in md
    assert "OUTRANGE_TOKEN_BBB" not in md
    assert "DEEP_TOKEN_CCC" not in md
    assert md.strip().startswith("|")


def test_sheet_to_markdown_exports_small_sheet_in_one_shot(tmp_path: Path) -> None:
    book = _write_ledger(tmp_path / "full.xlsx")

    result = sheet_to_markdown(str(book), sheet="网点清单")

    assert result["ok"] is True
    assert "容桂" in result["markdown"]
    assert "顺德" in result["markdown"]
    assert "网点" in result["markdown"]


def test_describe_and_markdown_deny_outside_path_without_leaking_name(
    tmp_path: Path, tmp_path_factory: pytest.TempPathFactory
) -> None:
    secret_dir = tmp_path_factory.mktemp("outside-describe")
    secret = secret_dir / "secret_ledger_outside.xlsx"
    wb = Workbook()
    wb.active.title = "SecretSheet"
    wb.active.append(["hidden_col"])
    wb.active.append(["SHOULD_NOT_LEAK"])
    wb.save(secret)

    described = describe_workbook(str(secret))
    assert described["ok"] is False
    assert described["issue"] == "path_denied"
    assert described.get("sheets") in ([], None) or described["sheets"] == []
    assert "secret_ledger_outside" not in described["message"]
    assert "SHOULD_NOT_LEAK" not in described["message"]
    assert "SecretSheet" not in described["message"]

    markdown = sheet_to_markdown(str(secret), start_row=1, end_row=2)
    assert markdown["ok"] is False
    assert markdown["issue"] == "path_denied"
    assert not (markdown.get("markdown") or "").strip()
    assert "secret_ledger_outside" not in markdown["message"]
    assert "SHOULD_NOT_LEAK" not in markdown["message"]


def test_describe_workbook_missing_file(tmp_path: Path) -> None:
    missing = tmp_path / "absent.xlsx"
    result = describe_workbook(str(missing))
    assert result["ok"] is False
    assert result["issue"] == "missing"
    assert "不存在" in result["message"] or "not found" in result["message"].lower()


def test_describe_csv_uses_header_and_preview(tmp_path: Path) -> None:
    csv_path = tmp_path / "branch.csv"
    csv_path.write_text(
        "网点,城市\n容桂,顺德\n李沧,青岛\n",
        encoding="utf-8",
    )

    result = describe_workbook(str(csv_path), preview_rows=1)

    assert result["ok"] is True
    sheet = result["sheets"][0]
    assert sheet["name"] == "branch"
    assert sheet["headers"] == ["网点", "城市"]
    assert sheet["row_count"] == 3
    blob = " ".join(" ".join(row) for row in sheet["preview"])
    assert "容桂" in blob
    assert "李沧" not in blob


def test_sheet_to_markdown_csv_range(tmp_path: Path) -> None:
    csv_path = tmp_path / "tx.csv"
    csv_path.write_text(
        "id,note\n1,INRANGE_TOKEN_AAA\n2,OUTRANGE_TOKEN_BBB\n",
        encoding="utf-8",
    )

    result = sheet_to_markdown(str(csv_path), start_row=2, end_row=2)

    assert result["ok"] is True
    assert "INRANGE_TOKEN_AAA" in result["markdown"]
    assert "id" in result["markdown"]
    assert "OUTRANGE_TOKEN_BBB" not in result["markdown"]
