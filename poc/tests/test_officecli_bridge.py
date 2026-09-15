# -*- coding: utf-8 -*-
"""Real tests for excel_guard_mcp.officecli_bridge (skipped without the CLI)."""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from poc.excel_guard_mcp.officecli_bridge import (
    edit_workbook,
    inspect_workbook,
    render_workbook,
    validate_workbook,
)

pytestmark = pytest.mark.skipif(
    shutil.which("officecli") is None, reason="officecli binary not installed"
)


@pytest.fixture(autouse=True)
def _sandbox_workspace(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Confine path checks to the per-test temp dir."""
    monkeypatch.setenv("POC_WORKSPACE", str(tmp_path))


@pytest.fixture()
def workbook(tmp_path: Path) -> Path:
    path = tmp_path / "sales.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Region"
    ws["B1"] = "Sales"
    for row in (("华东", 100), ("华南", 200), ("华北", 300)):
        ws.append(row)
    wb.save(path)
    return path


def test_edit_writes_values_and_formula_autoevals(workbook: Path) -> None:
    result = edit_workbook(
        str(workbook),
        [
            {"command": "set", "path": "/Sheet1/B5", "props": {"value": "=SUM(B2:B4)"}},
            {"command": "set", "path": "/Sheet1/C1", "props": {"value": "备注"}},
        ],
    )

    assert result["ok"] is True, result
    assert result["summary"]["succeeded"] == 2
    assert result["validate"] == "Validation passed: no errors found."
    assert result["audit"]["ok"] is True

    ws = load_workbook(workbook, data_only=True)["Sheet1"]
    assert ws["B5"].value == 600  # officecli evaluated the formula on write
    assert ws["C1"].value == "备注"


def test_edit_rolls_back_atomically_on_bad_item(workbook: Path) -> None:
    result = edit_workbook(
        str(workbook),
        [
            {"command": "set", "path": "/Sheet1/B5", "props": {"value": 1}},
            {"command": "set", "path": "/NoSheet/B1", "props": {"value": 2}},
        ],
    )

    # officecli batch is atomic: one bad item rolls everything back.
    assert result["ok"] is False
    assert result["summary"].get("atomicRolledBack") is True
    assert load_workbook(workbook)["Sheet1"]["B5"].value is None


def test_edit_rejects_disallowed_verbs(workbook: Path) -> None:
    result = edit_workbook(
        str(workbook),
        [{"command": "raw-set", "path": "/Sheet1/B1", "props": {}}],
    )

    assert result["ok"] is False
    assert result["issue"] == "verb_not_allowed"
    # Nothing was written.
    assert load_workbook(workbook)["Sheet1"]["C1"].value is None


def test_edit_denies_path_outside_workspace(tmp_path: Path, monkeypatch) -> None:
    outside = tmp_path.parent / "outside.xlsx"
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(outside)

    result = edit_workbook(
        str(outside), [{"command": "set", "path": "/Sheet1/A1", "props": {"value": 1}}]
    )

    assert result["ok"] is False
    assert result["issue"] == "path_denied"


def test_edit_rejects_corrupt_file(tmp_path: Path) -> None:
    bad = tmp_path / "corrupt.xlsx"
    bad.write_bytes(b"not a zip")

    result = edit_workbook(
        str(bad), [{"command": "set", "path": "/Sheet1/A1", "props": {"value": 1}}]
    )

    assert result["ok"] is False
    assert result["issue"] == "corrupt"


def test_inspect_returns_cell_format(workbook: Path) -> None:
    result = inspect_workbook(str(workbook), "/Sheet1/B2", depth=0)

    assert result["ok"] is True, result
    node = result["node"]["results"][0]
    assert node["path"] == "/Sheet1/B2"
    assert node["text"] == "100"


def test_validate_workbook_passes(workbook: Path) -> None:
    result = validate_workbook(str(workbook))

    assert result["ok"] is True
    assert "校验通过" in result["message"]


def test_render_writes_png_inside_workspace(workbook: Path, tmp_path: Path) -> None:
    out = tmp_path / "render" / "sales.png"

    result = render_workbook(str(workbook), str(out))

    assert result["ok"] is True, result
    produced = Path(result["screenshot"])
    assert produced.is_relative_to(tmp_path)
    assert zipfile.is_zipfile(workbook)  # render must not corrupt the source
    # officecli renders through a headless browser; the PNG is non-trivial.
    assert produced.stat().st_size > 0
