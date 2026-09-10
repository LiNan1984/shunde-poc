# -*- coding: utf-8 -*-
"""Integration tests for the FastMCP ``report-visualizer`` server.

Mirrors ``test_mcp_server.py``: exercises the real FastMCP tool registry and
invokes every tool through ``mcp.call_tool`` (no stdio transport). The full
demo pipeline is driven end to end through the MCP layer: pivot -> 5 charts
-> docx assembly, then the docx is reopened and structurally verified.
"""

from __future__ import annotations

import asyncio
import json
import os
import subprocess
import sys
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn

from poc.report_mcp.server import main, mcp

EXPECTED_TOOLS = {
    "pivot_table_tool",
    "render_bar_tool",
    "render_line_tool",
    "render_pie_tool",
    "render_scatter_tool",
    "render_heatmap_tool",
    "render_docx_report_tool",
}
REPO_ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(autouse=True)
def _sandbox_workspace(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("POC_WORKSPACE", str(tmp_path))


ROWS = [
    {"region": "北区", "quarter": "Q1", "product": "储蓄", "sales": 100, "customers": 30},
    {"region": "北区", "quarter": "Q2", "product": "信贷", "sales": 140, "customers": 34},
    {"region": "南区", "quarter": "Q1", "product": "储蓄", "sales": 80, "customers": 20},
    {"region": "南区", "quarter": "Q2", "product": "理财", "sales": 120, "customers": 28},
    {"region": "东区", "quarter": "Q1", "product": "信贷", "sales": 60, "customers": 18},
    {"region": "东区", "quarter": "Q2", "product": "储蓄", "sales": 90, "customers": 22},
]


def _list_tools() -> dict[str, Any]:
    tools = asyncio.run(mcp.list_tools())
    return {tool.name: tool for tool in tools}


def _call_tool(name: str, arguments: dict[str, Any]) -> dict[str, Any]:
    result = asyncio.run(mcp.call_tool(name, arguments))
    if isinstance(result, dict):  # forward-compatible structured return
        return result
    assert result, f"tool {name!r} returned no content blocks"
    text = getattr(result[0], "text", None)
    assert text is not None, f"tool {name!r} returned non-text block: {result[0]!r}"
    payload = json.loads(text)
    assert isinstance(payload, dict)
    return payload


# ---------------------------------------------------------------------------
# Registration
# ---------------------------------------------------------------------------


def test_server_registers_exactly_seven_tools() -> None:
    tools = _list_tools()

    assert len(tools) == 7
    assert set(tools) == EXPECTED_TOOLS


def test_server_name_is_report_visualizer() -> None:
    assert mcp.name == "report-visualizer"


@pytest.mark.parametrize("tool_name", sorted(EXPECTED_TOOLS))
def test_registered_tools_have_descriptions(tool_name: str) -> None:
    desc = _list_tools()[tool_name].description
    assert desc and desc.strip()


def test_chart_tools_schemas_accept_output_path(tmp_path: Path) -> None:
    for tool_name in EXPECTED_TOOLS - {"pivot_table_tool"}:
        schema = _list_tools()[tool_name].inputSchema
        assert schema["type"] == "object"
        assert "output_path" in schema["properties"]


# ---------------------------------------------------------------------------
# Every tool actually callable through the MCP registry
# ---------------------------------------------------------------------------


def test_mcp_pivot_table_tool() -> None:
    payload = _call_tool(
        "pivot_table_tool",
        {"data": ROWS, "rows": ["quarter"], "cols": ["region"], "value": "sales"},
    )
    assert payload["ok"] is True
    assert payload["row_count"] == 2
    assert payload["col_count"] == 3
    assert isinstance(payload["table"], list) and payload["table"]


def test_mcp_pivot_table_tool_error_shape() -> None:
    payload = _call_tool(
        "pivot_table_tool",
        {"data": [], "rows": ["quarter"], "cols": [], "value": "sales"},
    )
    assert payload["ok"] is False
    assert payload["table"] == []
    assert "empty" in payload["message"].lower()


def test_mcp_chart_tools_render_pngs(tmp_path: Path) -> None:
    calls = {
        "render_bar_tool": {
            "data": ROWS, "x": "quarter", "y": "sales", "series": "region",
            "output_path": str(tmp_path / "bar.png"),
        },
        "render_line_tool": {
            "data": ROWS, "x": "customers", "y": "sales", "series": "region",
            "output_path": str(tmp_path / "line.png"),
        },
        "render_pie_tool": {
            "data": ROWS, "label": "product", "value": "sales",
            "output_path": str(tmp_path / "pie.png"),
        },
        "render_scatter_tool": {
            "data": ROWS, "x": "customers", "y": "sales", "color": "region",
            "output_path": str(tmp_path / "scatter.png"),
        },
        "render_heatmap_tool": {
            "matrix": [[1, 2, 3], [4, 5, 6]],
            "x_labels": ["Q1", "Q2", "Q3"],
            "y_labels": ["北区", "南区"],
            "output_path": str(tmp_path / "heatmap.png"),
        },
    }

    for tool_name, arguments in calls.items():
        payload = _call_tool(tool_name, arguments)
        assert payload["ok"] is True, (tool_name, payload)
        assert Path(payload["output_path"]).is_file(), tool_name
        assert payload["chart_type"] in tool_name  # bar/line/pie/scatter/heatmap


def test_mcp_chart_tool_denies_path_outside_workspace() -> None:
    payload = _call_tool(
        "render_bar_tool",
        {"data": ROWS, "x": "region", "y": "sales", "output_path": "/etc/evil.png"},
    )
    assert payload["ok"] is False
    assert payload.get("issue") == "path_denied"


def test_mcp_full_pipeline_assembles_docx(tmp_path: Path) -> None:
    """Pivot + 5 charts + docx, all through MCP tools; reopen to verify."""
    pivot = _call_tool(
        "pivot_table_tool",
        {"data": ROWS, "rows": ["quarter"], "cols": ["region"], "value": "sales"},
    )
    assert pivot["ok"] is True

    chart_paths: list[str] = []
    chart_specs = [
        ("render_bar_tool", {"data": ROWS, "x": "region", "y": "sales"}),
        ("render_line_tool", {"data": ROWS, "x": "customers", "y": "sales"}),
        ("render_pie_tool", {"data": ROWS, "label": "product", "value": "sales"}),
        ("render_scatter_tool", {"data": ROWS, "x": "customers", "y": "sales"}),
        (
            "render_heatmap_tool",
            {"data": ROWS, "row": "region", "col": "quarter", "value": "sales"},
        ),
    ]
    for i, (tool_name, args) in enumerate(chart_specs):
        args["output_path"] = str(tmp_path / f"chart_{i}.png")
        payload = _call_tool(tool_name, args)
        assert payload["ok"] is True, (tool_name, payload)
        chart_paths.append(payload["output_path"])

    table_records = pivot["table"]
    headers = list(table_records[0].keys())
    docx_path = tmp_path / "mcp_report.docx"
    report = _call_tool(
        "render_docx_report_tool",
        {
            "title": "MCP 全链路报告",
            "header": "顺德农商行 POC",
            "footer": "演示页",
            "output_path": str(docx_path),
            "sections": [
                {
                    "heading": "季度 x 区域交叉表",
                    "level": 2,
                    "paragraphs": ["本表经 pivot_table_tool 生成。"],
                    "table": {
                        "headers": headers,
                        "rows": [list(r.values()) for r in table_records],
                    },
                    "charts": chart_paths,
                }
            ],
        },
    )
    assert report["ok"] is True, report
    assert report["charts_inserted"] == 5
    assert docx_path.is_file()

    doc = Document(str(docx_path))
    assert len(doc.tables) == 1
    drawings = doc.element.findall(".//" + qn("w:drawing"))
    assert len(drawings) == 5
    assert "MCP 全链路报告" in "\n".join(p.text for p in doc.paragraphs)


def test_mcp_docx_tool_validates_extension(tmp_path: Path) -> None:
    payload = _call_tool(
        "render_docx_report_tool",
        {"title": "t", "sections": [], "output_path": str(tmp_path / "x.txt")},
    )
    assert payload["ok"] is False
    assert ".docx" in payload["message"]


# ---------------------------------------------------------------------------
# Lifecycle
# ---------------------------------------------------------------------------


def test_main_is_callable_and_invokes_mcp_run(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[tuple[tuple, dict]] = []

    def _fake_run(*args: Any, **kwargs: Any) -> None:
        calls.append((args, kwargs))

    monkeypatch.setattr(mcp, "run", _fake_run)

    main()

    assert calls == [((), {})]


def test_importing_server_module_does_not_start_stdio() -> None:
    env = {**os.environ, "PYTHONPATH": str(REPO_ROOT)}
    completed = subprocess.run(
        [
            sys.executable,
            "-c",
            "from poc.report_mcp.server import mcp, main; "
            "assert mcp.name == 'report-visualizer'; "
            "assert callable(main); "
            "print('imported-ok')",
        ],
        cwd=str(REPO_ROOT),
        env=env,
        capture_output=True,
        text=True,
        timeout=20,
    )

    assert completed.returncode == 0, completed.stderr
    assert "imported-ok" in completed.stdout
